"""
OrgSight Feature Tracker Builder
Generates a formatted Excel workbook tracking all features, their status, complexity,
and remarks/feedback from the OrgSight DTS document.
"""

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from datetime import date

# ─── DATA ────────────────────────────────────────────────────────────────────

FEATURES = [
    {
        "id": 1,
        "category": "Core Modeling",
        "feature": "Org Modeling\n(Add / Close / Change / Clone Positions)",
        "description": (
            "Build a future-state org chart off the baseline in a draft. "
            "Actions are colour-coded and the as-is remains untouched. "
            "Includes add, close, change, and clone position operations."
        ),
        "reference": "Modeling 37:39–57:45",
        "status": "Completed",
        "complexity": "High",
        "built_summary": (
            "Full draft-mode modeling engine implemented. "
            "Position-level CRUD operations with colour-coded diff overlay on the org chart."
        ),
        "remarks": "",
    },
    {
        "id": 2,
        "category": "Core Modeling",
        "feature": "Rate-Card Costing",
        "description": (
            "Property combination (location + dept + grade) maps to indicative cost, "
            "auto-applied to new/changed positions. Quartile options and per-position "
            "manual override using existing FLC data."
        ),
        "reference": "Modeling 21:34–26:45",
        "status": "Completed",
        "complexity": "Medium",
        "built_summary": (
            "Rate-card lookup by location/dept/grade implemented. "
            "Quartile selector and manual override per position working. FLC data integrated."
        ),
        "remarks": "",
    },
    {
        "id": 3,
        "category": "Core Modeling",
        "feature": "Changes Report + Baseline vs To-Be View",
        "description": (
            "Size/cost delta (baseline / to-be / net) by any dimension, "
            "plus a side-by-side before/after comparison. "
            "This is the core client deliverable from any modeling exercise."
        ),
        "reference": "Modeling 1:01:42, 1:12:17",
        "status": "Completed",
        "complexity": "Medium",
        "built_summary": (
            "Delta report by dimension (headcount, cost) built. "
            "Side-by-side baseline vs to-be view rendered. Export to CSV/Excel supported."
        ),
        "remarks": "",
    },
    {
        "id": 4,
        "category": "Core Modeling",
        "feature": "Activity Analysis\n(People → Roles → Activities → Savings)",
        "description": (
            "Map activities to roles, apply automation/AI/stop/BPO levers with dates, "
            "and report FTE & cost impact over time. "
            "Aligns with AVA; per-person override capability."
        ),
        "reference": "Activity 11:50–49:50",
        "status": "Completed",
        "complexity": "High",
        "built_summary": (
            "Activity-to-role mapping with lever assignment (Automate / AI / Stop / BPO) built. "
            "FTE & cost savings timeline chart generated. Per-person override supported."
        ),
        "remarks": "",
    },
    {
        "id": 5,
        "category": "Core Modeling",
        "feature": "Action Dates + Phasing View",
        "description": (
            "Attach an effective date to each modeled change to show headcount "
            "removals/adds by month and compare full-year vs run-rate cost."
        ),
        "reference": "Modeling 1:05:30",
        "status": "Completed",
        "complexity": "Medium",
        "built_summary": (
            "Effective date field added to each change action. "
            "Monthly phasing waterfall and full-year vs run-rate cost comparison implemented."
        ),
        "remarks": "",
    },
    {
        "id": 6,
        "category": "Core Modeling",
        "feature": "Validate the To-Be State",
        "description": (
            "Flag modeling mistakes on the future state — e.g. closed manager with open reports, "
            "change with no reason code. Mirrors the existing Validate module but applied to drafts."
        ),
        "reference": "Modeling 40:39",
        "status": "Completed",
        "complexity": "Medium",
        "built_summary": (
            "Future-state validation rules implemented: orphaned reports, missing reason codes, "
            "span violations. Warnings surfaced in the draft before publishing."
        ),
        "remarks": "",
    },
    # ── PENDING ──────────────────────────────────────────────────────────────
    {
        "id": 7,
        "category": "Deepen Existing",
        "feature": "Calculated Properties\n(Formula Columns)",
        "description": (
            "Let users define a computed column off existing properties, "
            "e.g. cost-per-report, span ratio. "
            "Closes the Excel gap without needing a full Gizmo-style expression engine."
        ),
        "reference": "Deck (Expressions)",
        "status": "Pending",
        "complexity": "Medium",
        "built_summary": "",
        "remarks": "Formula parser + column renderer needed. Scope: basic arithmetic & IF logic.",
    },
    {
        "id": 8,
        "category": "Deepen Existing",
        "feature": "Properties Cleaning Heatmap",
        "description": (
            "Heatmap of missing values across critical fields, "
            "complementing existing Validate / Filter Errors modules. "
            "Highlights data quality gaps at a glance."
        ),
        "reference": "Deck (data validation)",
        "status": "Pending",
        "complexity": "Low",
        "built_summary": "",
        "remarks": "Can reuse Validate module data; primarily a UI/visualisation addition.",
    },
    {
        "id": 9,
        "category": "Deepen Existing",
        "feature": "Target-Span Alignment\n+ Micro-Team Detection",
        "description": (
            "Identify managers above/below a target span, quantify the FTE opportunity, "
            "and detect 1:1 / thin-layer structures. "
            "Extends the Spans & Layers module."
        ),
        "reference": "Modeling 1:22:52",
        "status": "Pending",
        "complexity": "Medium",
        "built_summary": "",
        "remarks": "Builds on S&L module. Needs configurable target-span input and FTE delta calc.",
    },
    {
        "id": 10,
        "category": "Beyond OrgVue",
        "feature": "Change-Impact Warnings\n(Leapfrog Detection)",
        "description": (
            "Auto-flag on the future state: too few managers, geography gaps, tenure skew. "
            "Proactive risk signals not currently in OrgVue."
        ),
        "reference": "Modeling 1:18:53",
        "status": "Pending",
        "complexity": "Medium",
        "built_summary": "",
        "remarks": "Differentiating feature vs OrgVue. Define warning thresholds as config.",
    },
    {
        "id": 11,
        "category": "Beyond OrgVue",
        "feature": "Editable-After-Kickoff\nModeling Config",
        "description": (
            "Allow modeling configuration (including rate card) to be changed "
            "after a modeling exercise has already started — "
            "fixing a known OrgVue limitation."
        ),
        "reference": "Modeling 30:57",
        "status": "Pending",
        "complexity": "Low",
        "built_summary": "",
        "remarks": "Requires versioning/re-apply logic for config changes mid-exercise.",
    },
]

# ─── STYLE HELPERS ───────────────────────────────────────────────────────────

def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border(sides="all"):
    s = Side(style="thin", color="C0C0C0")
    n = Side(style=None)
    if sides == "all":
        return Border(left=s, right=s, top=s, bottom=s)
    b = Border(
        left=s if "l" in sides else n,
        right=s if "r" in sides else n,
        top=s if "t" in sides else n,
        bottom=s if "b" in sides else n,
    )
    return b

def bold_border():
    s = Side(style="medium", color="888888")
    return Border(left=s, right=s, top=s, bottom=s)

FONT_BASE = "Calibri"

STATUS_COLORS = {
    "Completed": ("1E7145", "FFFFFF"),   # dark green bg, white text
    "Pending":   ("BF8F00", "FFFFFF"),   # amber bg, white text
    "In Progress": ("1F618D", "FFFFFF"), # blue bg, white text
}

COMPLEXITY_COLORS = {
    "High":   ("C0392B", "FFFFFF"),
    "Medium": ("E67E22", "FFFFFF"),
    "Low":    ("27AE60", "FFFFFF"),
}

CATEGORY_COLORS = {
    "Core Modeling":  "D6E4F0",
    "Deepen Existing": "FEF9E7",
    "Beyond OrgVue":  "F9EBEA",
}

HEADER_BG   = "1B2631"
HEADER_FG   = "FFFFFF"
TITLE_BG    = "154360"
ALT_ROW_BG  = "F2F2F2"
WHITE       = "FFFFFF"

# ─── WORKBOOK SETUP ──────────────────────────────────────────────────────────

wb = openpyxl.Workbook()

# ════════════════════════════════════════════════════════════════════════════
# SHEET 1 — TRACKER
# ════════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Feature Tracker"
ws.sheet_view.showGridLines = False

# ── Title row ────────────────────────────────────────────────────────────────
ws.merge_cells("A1:L1")
title_cell = ws["A1"]
title_cell.value = "OrgSight  |  Feature Development Tracker"
title_cell.font = Font(name=FONT_BASE, size=18, bold=True, color=HEADER_FG)
title_cell.fill = solid(TITLE_BG)
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 36

ws.merge_cells("A2:L2")
sub_cell = ws["A2"]
sub_cell.value = (
    f"Generated: {date.today().strftime('%d %b %Y')}    |    "
    "Features 1–6: Completed    |    Features 7–11: Pending    |    "
    "Source: OrgSight Detailed Functionality_DTS"
)
sub_cell.font = Font(name=FONT_BASE, size=10, italic=True, color="AAAAAA")
sub_cell.fill = solid("0E2030")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 18

# ── Column headers (row 3) ────────────────────────────────────────────────────
HEADERS = [
    "#",
    "Category",
    "Feature / Change Request",
    "Description",
    "Reference\n(DTS)",
    "Status",
    "Complexity",
    "What Was Built",
    "Remarks /\nFeedback",
    "Owner",
    "Target Date",
    "Last Updated",
]

COL_WIDTHS = [5, 18, 30, 52, 18, 14, 12, 45, 38, 16, 14, 14]

for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
    cell = ws.cell(row=3, column=col_idx, value=header)
    cell.font = Font(name=FONT_BASE, size=10, bold=True, color=HEADER_FG)
    cell.fill = solid(HEADER_BG)
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    cell.border = thin_border()
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[3].height = 32

# ── Data rows ────────────────────────────────────────────────────────────────
DATA_START = 4
last_category = None
current_row = DATA_START
feat_row_count = 0  # alternating color counter

for feat in FEATURES:
    cat_bg = CATEGORY_COLORS.get(feat["category"], WHITE)

    # Category separator label
    if feat["category"] != last_category:
        if last_category is not None:
            # spacer row between categories
            ws.row_dimensions[current_row].height = 6
            for c in range(1, 13):
                ws.cell(row=current_row, column=c).fill = solid("E8E8E8")
            current_row += 1

        # Category header row
        ws.merge_cells(f"A{current_row}:L{current_row}")
        cat_cell = ws.cell(row=current_row, column=1)
        cat_cell.value = f"  \u258c {feat['category'].upper()}"
        cat_cell.font = Font(name=FONT_BASE, size=9, bold=True, color="444444")
        cat_cell.fill = solid(cat_bg)
        cat_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 20
        current_row += 1
        last_category = feat["category"]
        feat_row_count = 0

    is_alt = feat_row_count % 2 == 1
    feat_row_count += 1
    row = current_row

    ws.row_dimensions[row].height = 72

    vals = [
        feat["id"],
        feat["category"],
        feat["feature"],
        feat["description"],
        feat["reference"],
        feat["status"],
        feat["complexity"],
        feat["built_summary"],
        feat["remarks"],
        "",            # Owner
        "",            # Target Date
        "",            # Last Updated
    ]

    for col_idx, val in enumerate(vals, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = Font(name=FONT_BASE, size=9, color="1A1A1A")
        cell.fill = solid(cat_bg)
        cell.alignment = Alignment(
            horizontal="left" if col_idx > 2 else "center",
            vertical="top",
            wrap_text=True,
        )
        cell.border = thin_border()

        # Styled pill-like coloring for Status (col 6) and Complexity (col 7)
        if col_idx == 6 and val in STATUS_COLORS:
            bg, fg = STATUS_COLORS[val]
            cell.fill = solid(bg)
            cell.font = Font(name=FONT_BASE, size=9, bold=True, color=fg)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        if col_idx == 7 and val in COMPLEXITY_COLORS:
            bg, fg = COMPLEXITY_COLORS[val]
            cell.fill = solid(bg)
            cell.font = Font(name=FONT_BASE, size=9, bold=True, color=fg)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        if col_idx == 1:
            cell.font = Font(name=FONT_BASE, size=10, bold=True, color="1A1A1A")

    current_row += 1

# ── Freeze panes & filters ───────────────────────────────────────────────────
ws.freeze_panes = "A4"
ws.auto_filter.ref = f"A3:L{current_row}"

# ── Data validation for Status column ────────────────────────────────────────
dv_status = DataValidation(
    type="list",
    formula1='"Completed,In Progress,Pending,On Hold,Cancelled"',
    allow_blank=True,
    showDropDown=False,
)
dv_status.sqref = f"F{DATA_START}:F{DATA_START + len(FEATURES) + 10}"
ws.add_data_validation(dv_status)

dv_complexity = DataValidation(
    type="list",
    formula1='"High,Medium,Low"',
    allow_blank=True,
    showDropDown=False,
)
dv_complexity.sqref = f"G{DATA_START}:G{DATA_START + len(FEATURES) + 10}"
ws.add_data_validation(dv_complexity)

# ════════════════════════════════════════════════════════════════════════════
# SHEET 2 — SUMMARY DASHBOARD
# ════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Summary Dashboard")
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:H1")
t = ws2["A1"]
t.value = "OrgSight  |  Summary Dashboard"
t.font = Font(name=FONT_BASE, size=16, bold=True, color=HEADER_FG)
t.fill = solid(TITLE_BG)
t.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 34

ws2.merge_cells("A2:H2")
s = ws2["A2"]
s.value = f"As of {date.today().strftime('%d %b %Y')}"
s.font = Font(name=FONT_BASE, size=9, italic=True, color="AAAAAA")
s.fill = solid("0E2030")
s.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[2].height = 16

# KPI cards
kpis = [
    ("Total Features", str(len(FEATURES)), "1B2631", "FFFFFF"),
    ("Completed", str(sum(1 for f in FEATURES if f["status"] == "Completed")), "1E7145", "FFFFFF"),
    ("Pending", str(sum(1 for f in FEATURES if f["status"] == "Pending")), "BF8F00", "FFFFFF"),
    ("High Complexity", str(sum(1 for f in FEATURES if f["complexity"] == "High")), "C0392B", "FFFFFF"),
    ("Medium Complexity", str(sum(1 for f in FEATURES if f["complexity"] == "Medium")), "E67E22", "FFFFFF"),
    ("Low Complexity", str(sum(1 for f in FEATURES if f["complexity"] == "Low")), "27AE60", "FFFFFF"),
]

for i, (label, value, bg, fg) in enumerate(kpis):
    col = 1 + i * 1
    ws2.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col)
    ws2.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col)

    label_cell = ws2.cell(row=4, column=col, value=label)
    label_cell.font = Font(name=FONT_BASE, size=9, color="888888")
    label_cell.fill = solid("F5F5F5")
    label_cell.alignment = Alignment(horizontal="center", vertical="center")
    label_cell.border = thin_border()
    ws2.row_dimensions[4].height = 18

    val_cell = ws2.cell(row=5, column=col, value=value)
    val_cell.font = Font(name=FONT_BASE, size=22, bold=True, color=fg)
    val_cell.fill = solid(bg)
    val_cell.alignment = Alignment(horizontal="center", vertical="center")
    val_cell.border = thin_border()
    ws2.row_dimensions[5].height = 42

for i in range(1, 7):
    ws2.column_dimensions[get_column_letter(i)].width = 18

# Feature breakdown table
ws2.row_dimensions[7].height = 14
hdr_row = 8
ws2.row_dimensions[hdr_row].height = 24
breakdown_headers = ["#", "Feature", "Category", "Status", "Complexity", "Reference"]
breakdown_widths  = [5, 38, 20, 14, 14, 18]
for ci, (h, w) in enumerate(zip(breakdown_headers, breakdown_widths), 1):
    c = ws2.cell(row=hdr_row, column=ci, value=h)
    c.font = Font(name=FONT_BASE, size=10, bold=True, color=HEADER_FG)
    c.fill = solid(HEADER_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border()
    ws2.column_dimensions[get_column_letter(ci)].width = w

for ri, feat in enumerate(FEATURES, start=1):
    r = hdr_row + ri
    ws2.row_dimensions[r].height = 20
    row_bg = CATEGORY_COLORS.get(feat["category"], WHITE)
    row_data = [
        feat["id"],
        feat["feature"].replace("\n", " "),
        feat["category"],
        feat["status"],
        feat["complexity"],
        feat["reference"],
    ]
    for ci, val in enumerate(row_data, 1):
        c = ws2.cell(row=r, column=ci, value=val)
        c.font = Font(name=FONT_BASE, size=9, color="1A1A1A")
        c.fill = solid(row_bg)
        c.alignment = Alignment(horizontal="left" if ci > 1 else "center", vertical="center")
        c.border = thin_border()

        if ci == 4 and val in STATUS_COLORS:
            bg, fg = STATUS_COLORS[val]
            c.fill = solid(bg)
            c.font = Font(name=FONT_BASE, size=9, bold=True, color=fg)
            c.alignment = Alignment(horizontal="center", vertical="center")

        if ci == 5 and val in COMPLEXITY_COLORS:
            bg, fg = COMPLEXITY_COLORS[val]
            c.fill = solid(bg)
            c.font = Font(name=FONT_BASE, size=9, bold=True, color=fg)
            c.alignment = Alignment(horizontal="center", vertical="center")

# ════════════════════════════════════════════════════════════════════════════
# SHEET 3 — CHANGE LOG
# ════════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Change Log")
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:G1")
t3 = ws3["A1"]
t3.value = "OrgSight  |  Change Log"
t3.font = Font(name=FONT_BASE, size=16, bold=True, color=HEADER_FG)
t3.fill = solid(TITLE_BG)
t3.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 34

ws3.merge_cells("A2:G2")
s3 = ws3["A2"]
s3.value = "Log every significant change, decision, or feedback here."
s3.font = Font(name=FONT_BASE, size=9, italic=True, color="AAAAAA")
s3.fill = solid("0E2030")
s3.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[2].height = 16

ws3.row_dimensions[3].height = 8

log_headers = ["Date", "Feature #", "Feature Name", "Change / Decision", "Raised By", "Status", "Notes"]
log_widths   = [14,     10,          30,               55,                  18,          16,       35]

for ci, (h, w) in enumerate(zip(log_headers, log_widths), 1):
    c = ws3.cell(row=4, column=ci, value=h)
    c.font = Font(name=FONT_BASE, size=10, bold=True, color=HEADER_FG)
    c.fill = solid(HEADER_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border()
    ws3.column_dimensions[get_column_letter(ci)].width = w
ws3.row_dimensions[4].height = 24

# Pre-fill a few sample rows
sample_logs = [
    (date.today().strftime("%d %b %Y"), "1–6", "All Core Features", "Initial build complete across agents 1–6", "Dev Team", "Closed", "Verified end-to-end"),
    (date.today().strftime("%d %b %Y"), "7",   "Calculated Properties", "Scoped: arithmetic + IF logic only (no full Gizmo engine)", "Product", "Open",   "Awaiting sprint assignment"),
    (date.today().strftime("%d %b %Y"), "8",   "Properties Heatmap",   "Reuse existing Validate data; UI work only",               "Product", "Open",   "Low effort, quick win"),
    (date.today().strftime("%d %b %Y"), "9",   "Target-Span Alignment","Configurable target-span input required",                   "Product", "Open",   "Extend S&L module"),
    (date.today().strftime("%d %b %Y"), "10",  "Change-Impact Warnings","Define threshold config before build",                     "Product", "Open",   "Differentiating vs OrgVue"),
    (date.today().strftime("%d %b %Y"), "11",  "Editable Config",      "Versioning / re-apply logic needed",                       "Product", "Open",   "Fixes known OrgVue gap"),
]

for ri, log in enumerate(sample_logs, start=5):
    ws3.row_dimensions[ri].height = 22
    is_alt = ri % 2 == 0
    for ci, val in enumerate(log, 1):
        c = ws3.cell(row=ri, column=ci, value=val)
        c.font = Font(name=FONT_BASE, size=9, color="1A1A1A")
        c.fill = solid(ALT_ROW_BG if is_alt else WHITE)
        c.alignment = Alignment(horizontal="left" if ci > 2 else "center", vertical="center", wrap_text=True)
        c.border = thin_border()

ws3.freeze_panes = "A5"

# ─── SAVE ────────────────────────────────────────────────────────────────────
OUTPUT_PATH = r"c:\Users\jagritjain\OneDrive - Alvarez and Marsal\Documents\APPS\ORG_ANALYSIS\OrgSight_Feature_Tracker.xlsx"
wb.save(OUTPUT_PATH)
print(f"Saved: {OUTPUT_PATH}")
