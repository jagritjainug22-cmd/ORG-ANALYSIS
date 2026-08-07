"""
Project-scoped lifecycle router (Phase 3).

Every endpoint lives under /projects/{project_id}/… and is protected by
require_project_access(), which enforces JWT auth, project membership,
deadline, and status checks in a single dependency.  project_id is always
sourced from the URL path -- never from a request body.
"""

from __future__ import annotations

import json
import logging
import os
import subprocess
import tempfile
from io import BytesIO
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

import numpy as np
import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import ORJSONResponse, StreamingResponse
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pydantic import BaseModel

from dependencies.auth import (
    get_current_user,
    require_project_access,
    require_dataset_lock_holder,
    require_dataset_lock_holder_for_dataset,
)
from services import dataset_lock_service
from services import duckdb_manager
from services.audit_service import write_audit_log
from services import db_service
from services.cleanup_service import apply_exclusion_filter, build_country_flag
from services.completeness_service import get_completeness_matrix
from services.crosstab_service import apply_others_grouping, generate_crosstab, generate_preview_data
from services.export_service import export_excel
from services.filter_error_service import filter_errors
from services.hierarchy_service import compute_avg_flc, compute_chains, compute_direct_span, compute_levels, compute_total_reports
from services.logging_service import get_activity_logs, get_user_stats, write_activity_log
from services.orgchart_render_service import render_scenario_svg, get_tree_structure
from services.orgchart_service import build_org_tree, build_tree_preserve_ancestors, make_json_serializable
from services.formula_service import apply_formulas_to_records, evaluate_formula, validate_expression as formula_validate_expression
from services.spans_layers_service import (
    get_insights,
    get_layer_employees,
    get_manager_detail,
    span_threshold,
    spans_and_layers,
)
from services.upload_service import read_excel_file, smart_read_excel
from services.validation_service import validate_org_data, validate_scenario
from services.column_mapping_service import auto_map_columns, auto_map_columns_with_feedback
from services.rationalisation_service import (
    rationalise,
    apply_rationalisation,
    persist_approved_mappings,
    list_learned_taxonomy_entries,
    update_learned_taxonomy_entry,
)

Image.MAX_IMAGE_PIXELS = None

router = APIRouter(prefix="/projects/{project_id}", tags=["lifecycle"])

# ---------------------------------------------------------------------------
# Pydantic request models
# ---------------------------------------------------------------------------

class SaveBaselineBody(BaseModel):
    name: str
    records: List[Dict]
    emp_col: str
    mgr_col: str
    fte_col: Optional[str] = None
    flc_col: Optional[str] = None
    job_title_col: Optional[str] = None
    country_col: Optional[str] = None
    func_col: Optional[str] = None
    subfunc_col: Optional[str] = None
    grade_col: Optional[str] = None
    division_col: Optional[str] = None
    entity_col: Optional[str] = None
    start_date_col: Optional[str] = None
    basic_pay_col: Optional[str] = None
    contract_type_col: Optional[str] = None
    status_col: Optional[str] = None


class ColumnConfigBody(BaseModel):
    emp_col: Optional[str] = None
    mgr_col: Optional[str] = None
    fte_col: Optional[str] = None
    flc_col: Optional[str] = None
    job_title_col: Optional[str] = None
    country_col: Optional[str] = None
    func_col: Optional[str] = None
    subfunc_col: Optional[str] = None
    grade_col: Optional[str] = None
    division_col: Optional[str] = None
    entity_col: Optional[str] = None
    start_date_col: Optional[str] = None
    basic_pay_col: Optional[str] = None
    contract_type_col: Optional[str] = None
    status_col: Optional[str] = None

class MoveBody(BaseModel):
    emp_id: str
    new_mgr_id: Optional[str] = None
    effective_date: Optional[str] = None

class EditBody(BaseModel):
    emp_id: str
    updates: Dict
    effective_date: Optional[str] = None

class AddBody(BaseModel):
    record: Dict
    emp_id: str
    mgr_id: Optional[str] = None
    level: Optional[int] = None
    fte: Optional[float] = None
    flc: Optional[float] = None
    effective_date: Optional[str] = None

class FlagBody(BaseModel):
    emp_id: str
    flagged: bool = True
    effective_date: Optional[str] = None

class CloneBody(BaseModel):
    source_emp_id: str
    new_emp_id: str
    new_mgr_id: Optional[str] = None
    effective_date: Optional[str] = None

class BulkDateBody(BaseModel):
    change_ids: List[int]
    effective_date: Optional[str] = None

class BulkFlagBody(BaseModel):
    emp_ids: List[str]
    flagged: bool = True
    effective_date: Optional[str] = None

class BulkEditPropertyBody(BaseModel):
    emp_ids: List[str]
    field: str
    value: Any
    effective_date: Optional[str] = None

class BulkMoveBody(BaseModel):
    emp_ids: List[str]
    new_mgr_id: Optional[str] = None
    effective_date: Optional[str] = None

class FormulaCreateBody(BaseModel):
    col_name: str
    expression: str

class FormulaPreviewBody(BaseModel):
    expression: str
    sample_size: int = 5
    data: List[Dict]
    available_columns: Optional[List[str]] = None

class CompletenessHeatmapBody(BaseModel):
    data: List[Dict]
    fields: List[str]
    group_col: str

class ScenarioBody(BaseModel):
    name: str
    description: Optional[str] = ""
    source_scenario_id: Optional[int] = None
    rate_card_id: Optional[int] = None
    rate_card_quartile: Optional[str] = "p50"


class RateCardGenerateBody(BaseModel):
    name: str
    property_cols: List[str]
    cost_col: Optional[str] = None
    min_sample: int = 3


class RateCardPreviewBody(BaseModel):
    property_cols: List[str]
    cost_col: Optional[str] = None
    min_sample: int = 3


class RateCardRowBody(BaseModel):
    composite_key: str
    p25: Optional[float] = None
    p50: Optional[float] = None
    p75: Optional[float] = None
    avg_cost: Optional[float] = None


class ScenarioRateCardBody(BaseModel):
    rate_card_id: Optional[int] = None
    rate_card_quartile: str = "p50"


class RateCardLookupBody(BaseModel):
    values: Dict[str, Any]


class ColumnMappingBody(BaseModel):
    columns: List[str]
    sample_rows: List[Dict]


class RationaliseBody(BaseModel):
    records: List[Dict]
    func_col: Optional[str] = None
    subfunc_col: Optional[str] = None
    title_col: Optional[str] = None
    ignore_learned_aliases: bool = True


class RationaliseApplyBody(BaseModel):
    records: List[Dict]
    func_col: Optional[str] = None
    subfunc_col: Optional[str] = None
    title_col: Optional[str] = None
    approved_functions: List[Dict] = []
    approved_subfunctions: List[Dict] = []
    approved_titles: List[Dict] = []
    dataset_id: Optional[int] = None
    enrich_learned: bool = True


class LearnedTaxonomyPatchBody(BaseModel):
    entry_id: str
    resolved: Optional[str] = None
    disabled: Optional[bool] = None
    apply_to_matching: bool = False
    delete: bool = False


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _require_dataset_in_project(dataset_id: int, project_id: int) -> Dict[str, Any]:
    """Load a dataset and verify it belongs to the given project."""
    ds = db_service.get_dataset(dataset_id)
    if not ds or ds.get("project_id") != project_id:
        raise HTTPException(status_code=404, detail="Dataset not found in this project")
    return ds


def _require_scenario_in_project(scenario_id: int, project_id: int) -> tuple[Dict[str, Any], Dict[str, Any]]:
    """Load a scenario + its parent dataset, verifying project ownership.
    Returns (scenario, dataset)."""
    scenario = db_service.get_scenario(scenario_id)
    if not scenario:
        raise HTTPException(status_code=404, detail="Scenario not found")
    ds = db_service.get_dataset(scenario["dataset_id"])
    if not ds or ds.get("project_id") != project_id:
        raise HTTPException(status_code=404, detail="Scenario not found in this project")
    return scenario, ds


# ---------------------------------------------------------------------------
# Inkscape-less PPT / PDF fallbacks (moved from main.py)
# ---------------------------------------------------------------------------

def _scenario_summary_rows(scenario: dict, dataset: dict, records: list) -> dict:
    summary = db_service.get_scenario_summary(scenario["id"])
    change_log = db_service.get_change_log(scenario["id"])
    return {
        "title": f"OrgSight 2.0 -- {dataset['name']} -- {scenario['name']}",
        "summary": summary,
        "change_log": change_log,
        "record_count": len(records),
    }


def _render_scenario_summary_pdf(scenario, dataset, records) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    data = _scenario_summary_rows(scenario, dataset, records)
    s = data["summary"]
    cur = s.get("current", {})
    base = s.get("baseline", {})
    delta = s.get("delta", {})
    flagged = s.get("flagged_removed", {})

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=1.5*cm, rightMargin=1.5*cm,
        topMargin=1.5*cm, bottomMargin=1.5*cm,
        title=data["title"],
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=colors.HexColor("#0b234b"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=colors.HexColor("#0b234b"))
    body = styles["BodyText"]

    flow = []
    flow.append(Paragraph(data["title"], h1))
    flow.append(Spacer(1, 0.3*cm))
    flow.append(Paragraph(
        f"Dataset #{dataset['id']} &middot; {data['record_count']:,} records", body,
    ))
    flow.append(Spacer(1, 0.5*cm))
    flow.append(Paragraph("Headline metrics", h2))

    metric_rows = [
        ["", "Baseline", "Current", "Delta"],
        ["Headcount", f"{base.get('headcount', 0):,}", f"{cur.get('headcount', 0):,}", f"{delta.get('headcount', 0):+,}"],
        ["Total FTE", f"{base.get('total_fte', 0):,.1f}", f"{cur.get('total_fte', 0):,.1f}", f"{delta.get('fte', 0):+,.1f}"],
        ["Total Cost", f"${base.get('total_cost', 0):,.0f}", f"${cur.get('total_cost', 0):,.0f}", f"${delta.get('cost', 0):+,.0f}"],
        ["Flagged removed", "0", f"{flagged.get('count', 0):,}", f"-{flagged.get('count', 0):,}"],
    ]
    t = Table(metric_rows, colWidths=[4*cm, 4*cm, 4*cm, 4*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b234b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d6dee8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fb")]),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
    ]))
    flow.append(t)
    flow.append(Spacer(1, 0.8*cm))

    flow.append(Paragraph("Change log", h2))
    change_rows = [["When", "Action", "Employee", "User", "Details"]]
    for entry in data["change_log"]:
        when = (entry.get("created_at") or "")[:19].replace("T", " ")
        details = entry.get("details_json") or ""
        if len(details) > 80:
            details = details[:80] + "..."
        change_rows.append([
            when, entry.get("action", ""), entry.get("emp_id", "") or "",
            entry.get("username", "") or "", details,
        ])
    if len(change_rows) == 1:
        change_rows.append(["--", "no changes yet", "", "", ""])
    log_table = Table(change_rows, colWidths=[3.5*cm, 2.0*cm, 3.0*cm, 2.5*cm, 6.0*cm], repeatRows=1)
    log_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b234b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d6dee8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fb")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    flow.append(log_table)
    flow.append(Spacer(1, 0.5*cm))
    flow.append(Paragraph(
        "<i>Note: this is the text summary export. For the full visual chart "
        "with editable shapes, install Inkscape on the server.</i>", body,
    ))
    doc.build(flow)
    return buf.getvalue()


def _svg_to_png(svg_content: str, max_width: int = 3840) -> bytes:
    """Convert an SVG string to high-res PNG bytes.

    Pipeline: SVG → svglib Drawing → reportlab PDF → pypdfium2 PNG.
    This avoids a hard dependency on the native Cairo DLL (libcairo-2.dll)
    which is difficult to provision on Windows.
    """
    import tempfile, os
    from svglib.svglib import svg2rlg
    from reportlab.graphics import renderPDF
    import pypdfium2 as pdfium

    with tempfile.NamedTemporaryFile(
        suffix=".svg", delete=False, mode="w", encoding="utf-8",
    ) as f:
        f.write(svg_content)
        svg_path = f.name

    try:
        drawing = svg2rlg(svg_path)
        if drawing is None:
            raise ValueError("svglib failed to parse SVG")

        pdf_buf = BytesIO()
        renderPDF.drawToFile(drawing, pdf_buf)
        pdf_buf.seek(0)

        pdf_doc = pdfium.PdfDocument(pdf_buf.getvalue())
        page = pdf_doc[0]
        scale = max_width / max(drawing.width, 1)
        bitmap = page.render(scale=scale)
        pil_image = bitmap.to_pil()

        png_buf = BytesIO()
        pil_image.save(png_buf, format="PNG")
        return png_buf.getvalue()
    finally:
        os.unlink(svg_path)


def _svg_to_pdf_bytes(svg_content: str) -> bytes:
    """Convert an SVG string to PDF bytes via svglib + reportlab."""
    import tempfile, os
    from svglib.svglib import svg2rlg
    from reportlab.graphics import renderPDF

    with tempfile.NamedTemporaryFile(
        suffix=".svg", delete=False, mode="w", encoding="utf-8",
    ) as f:
        f.write(svg_content)
        svg_path = f.name

    try:
        drawing = svg2rlg(svg_path)
        if drawing is None:
            raise ValueError("svglib failed to parse SVG")
        pdf_buf = BytesIO()
        renderPDF.drawToFile(drawing, pdf_buf)
        return pdf_buf.getvalue()
    finally:
        os.unlink(svg_path)


def _add_pptx_title_slide(prs, dataset, scenario) -> None:
    from pptx.util import Inches as I, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from datetime import datetime

    C_NAVY = RGBColor(0x01, 0x24, 0x4A)
    C_GOLD = RGBColor(0xC5, 0x96, 0x0C)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY
    bg.line.fill.background()

    tb = slide.shapes.add_textbox(I(1), I(2.0), I(11.3), I(1.2))
    p = tb.text_frame.paragraphs[0]
    p.text = "OrgSight 2.0"
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = C_WHITE

    sub = slide.shapes.add_textbox(I(1), I(3.2), I(11.3), I(0.6))
    p2 = sub.text_frame.paragraphs[0]
    p2.text = "Organizational Structure Analysis"
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.runs[0]
    run2.font.size = Pt(22)
    run2.font.color.rgb = C_GOLD

    line = slide.shapes.add_shape(1, I(5), I(4.0), I(3.3), Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = C_GOLD
    line.line.fill.background()

    info = slide.shapes.add_textbox(I(1), I(4.4), I(11.3), I(1.6))
    tf = info.text_frame
    for txt in [
        f"Scenario: {scenario['name']}",
        f"Dataset: {dataset['name']}",
        f"Exported: {datetime.now().strftime('%d %b %Y')}",
    ]:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = txt
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(14)
            r.font.color.rgb = C_WHITE

    foot = slide.shapes.add_textbox(I(7), I(6.8), I(6), I(0.4))
    fp = foot.text_frame.paragraphs[0]
    fp.text = "Confidential — Alvarez & Marsal"
    fp.alignment = PP_ALIGN.RIGHT
    for r in fp.runs:
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x8A, 0x9A, 0xB4)


def _add_pptx_summary_slide(prs, summary, dataset, scenario, page_num: int = 2) -> None:
    from pptx.util import Inches as I, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    C_NAVY = RGBColor(0x01, 0x24, 0x4A)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    C_GREY = RGBColor(0x4F, 0x60, 0x77)

    cur = summary.get("current", {})
    base = summary.get("baseline", {})
    delta = summary.get("delta", {})

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, I(0.7))
    band.fill.solid()
    band.fill.fore_color.rgb = C_NAVY
    band.line.fill.background()
    ttl = slide.shapes.add_textbox(I(0.5), I(0.12), I(12), I(0.5))
    tp = ttl.text_frame.paragraphs[0]
    tp.text = f"Executive Summary — {scenario['name']}"
    for r in tp.runs:
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = C_WHITE

    # 1. Comparison Table (Left Side)
    # Metric | Baseline | Target (To-Be) | Delta
    table_shape = slide.shapes.add_table(4, 4, I(0.5), I(1.5), I(7.0), I(2.5))
    tbl = table_shape.table
    tbl.columns[0].width = I(2.2)
    tbl.columns[1].width = I(1.6)
    tbl.columns[2].width = I(1.6)
    tbl.columns[3].width = I(1.6)

    headers = ["Metric", "Baseline", "Current (To-Be)", "Delta"]
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(11)
                r.font.color.rgb = C_WHITE

    metrics_rows = [
        ("Headcount", f"{base.get('headcount', 0):,}", f"{cur.get('headcount', 0):,}", f"{delta.get('headcount', 0):+d}"),
        ("Total FTE", f"{base.get('total_fte', 0):,.1f}", f"{cur.get('total_fte', 0):,.1f}", f"{delta.get('fte', 0):+.1f}"),
        ("Total Cost", f"${base.get('total_cost', 0):,.0f}", f"${cur.get('total_cost', 0):,.0f}", f"${delta.get('cost', 0):+,.0f}")
    ]

    for r_idx, (m, b_val, c_val, d_val) in enumerate(metrics_rows, start=1):
        vals = [m, b_val, c_val, d_val]
        for c_idx, v in enumerate(vals):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = v
            if r_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF5, 0xF7, 0xFA)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT
                for r in p.runs:
                    r.font.size = Pt(10)
                    if c_idx == 3:
                        r.font.bold = True
                        if "-" in d_val:
                            r.font.color.rgb = RGBColor(0x27, 0xAE, 0x60)
                        elif "+" in d_val:
                            r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

    # 2. Savings and Impact Panel (Right Side)
    panel_left = I(8.0)
    panel = slide.shapes.add_shape(1, panel_left, I(1.5), I(4.8), I(4.5))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(0xFA, 0xFB, 0xFC)
    panel.line.color.rgb = RGBColor(0xDC, 0xE4, 0xEE)
    
    # Title inside panel
    shdr = slide.shapes.add_shape(1, panel_left, I(1.5), I(4.8), I(0.45))
    shdr.fill.solid()
    shdr.fill.fore_color.rgb = C_NAVY
    shdr.line.fill.background()
    shdr_tb = slide.shapes.add_textbox(panel_left, I(1.55), I(4.8), I(0.4))
    shp = shdr_tb.text_frame.paragraphs[0]
    shp.text = "SAVINGS & IMPACT SUMMARY"
    shp.alignment = PP_ALIGN.CENTER
    for r in shp.runs:
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = C_WHITE

    # Fetch phasing to display financial savings
    phasing = db_service.get_phasing_view(scenario["id"], fy_start_month=1)
    ph_sum = phasing.get("summary", {})
    
    info_points = [
        ("Full-Year Annualized Savings:", f"${ph_sum.get('full_year_savings', 0):,.0f}"),
        ("Prorated In-Year Savings:", f"${ph_sum.get('in_year_savings', 0):,.0f}"),
        ("Months Remaining in FY:", f"{ph_sum.get('months_remaining_in_fy', 0)}"),
        ("Flagged Removed Roles:", f"{summary.get('flagged_removed', {}).get('count', 0)}"),
        ("Flagged Removed Value:", f"${summary.get('flagged_removed', {}).get('cost', 0):,.0f}"),
        ("Total Changes Logged:", f"{summary.get('change_count', 0)}"),
    ]
    
    tb_content = slide.shapes.add_textbox(panel_left + I(0.2), I(2.1), I(4.4), I(3.8))
    tf = tb_content.text_frame
    tf.word_wrap = True
    for idx, (lbl, val) in enumerate(info_points):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run_lbl = p.add_run()
        run_lbl.text = f"{lbl} "
        run_lbl.font.size = Pt(11)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = C_NAVY
        
        run_val = p.add_run()
        run_val.text = val
        run_val.font.size = Pt(11)
        run_val.font.bold = True
        if "Savings" in lbl:
            run_val.font.color.rgb = RGBColor(0x27, 0xAE, 0x60)
        else:
            run_val.font.color.rgb = C_GREY

    _add_pptx_footer(slide, prs, scenario["name"], page_num)


def _add_pptx_comparison_slide(prs, dataset_id: int, active_scenario_id: int, scenario_name: str, page_num: int) -> None:
    from pptx.util import Inches as I, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    C_NAVY = RGBColor(0x01, 0x24, 0x4A)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, I(0.7))
    band.fill.solid()
    band.fill.fore_color.rgb = C_NAVY
    band.line.fill.background()
    ttl = slide.shapes.add_textbox(I(0.5), I(0.12), I(12), I(0.5))
    tp = ttl.text_frame.paragraphs[0]
    tp.text = "Scenario Comparison Analysis"
    for r in tp.runs:
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = C_WHITE

    # Get comparison data
    scenarios_list = db_service.list_scenarios(dataset_id)
    headers = ["Scenario Name", "HC", "FTE", "Cost", "Δ HC", "Δ Cost", "Changes"]
    
    rows = len(scenarios_list) + 1
    cols = len(headers)
    table_shape = slide.shapes.add_table(rows, cols, I(0.5), I(1.2), I(12.33), I(min(5.0, 0.4 * rows)))
    tbl = table_shape.table
    
    tbl.columns[0].width = I(3.0)
    for c in range(1, cols):
        tbl.columns[c].width = I(1.55)

    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(11)
                r.font.color.rgb = C_WHITE

    for i, s in enumerate(scenarios_list, start=1):
        try:
            s_sum = db_service.get_scenario_summary(s["id"])
            hc = f"{s_sum['current']['headcount']:,}"
            fte = f"{s_sum['current']['total_fte']:,.1f}"
            cost = f"${s_sum['current']['total_cost']:,.0f}"
            dhc = f"{s_sum['delta']['headcount']:+d}"
            dcost = f"${s_sum['delta']['cost']:+,.0f}"
            ch = f"{s_sum['change_count']}"
            name = s["name"]
            if s["id"] == active_scenario_id:
                name += " (Active)"
            if s["is_promoted"]:
                name += " [Baseline]"
                
            vals = [name, hc, fte, cost, dhc, dcost, ch]
            for c, v in enumerate(vals):
                cell = tbl.cell(i, c)
                cell.text = v
                if i % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xF5, 0xF7, 0xFA)
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
                    for r in p.runs:
                        r.font.size = Pt(10)
                        if s["id"] == active_scenario_id:
                            r.font.bold = True
                            r.font.color.rgb = C_NAVY
        except Exception:
            continue

    _add_pptx_footer(slide, prs, scenario_name, page_num)


def _add_pptx_phasing_slide(prs, scenario_id: int, scenario_name: str, page_num: int) -> None:
    from pptx.util import Inches as I, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    C_NAVY = RGBColor(0x01, 0x24, 0x4A)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, I(0.7))
    band.fill.solid()
    band.fill.fore_color.rgb = C_NAVY
    band.line.fill.background()
    ttl = slide.shapes.add_textbox(I(0.5), I(0.12), I(12), I(0.5))
    tp = ttl.text_frame.paragraphs[0]
    tp.text = "Implementation Phasing Timeline"
    for r in tp.runs:
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = C_WHITE

    # Get phasing data
    phasing = db_service.get_phasing_view(scenario_id, fy_start_month=1)
    buckets = phasing.get("monthly", [])
    
    headers = ["Month", "Adds", "Removals", "Net HC", "Cost Delta", "Cum. Cost Delta"]
    rows = len(buckets) + 1
    cols = len(headers)
    table_shape = slide.shapes.add_table(rows, cols, I(0.5), I(1.2), I(12.33), I(min(5.0, 0.4 * rows)))
    tbl = table_shape.table

    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(11)
                r.font.color.rgb = C_WHITE

    for i, b in enumerate(buckets, start=1):
        month = b["label"]
        adds = f"{b['adds_count']}"
        rem = f"{b['removes_count']}"
        net_hc = f"{b['net_hc_delta']:+d}"
        c_delta = f"${b['cost_delta']:+,.0f}"
        cum_cost = f"${b['cumulative_cost_delta']:+,.0f}"
        
        vals = [month, adds, rem, net_hc, c_delta, cum_cost]
        for c, v in enumerate(vals):
            cell = tbl.cell(i, c)
            cell.text = v
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF5, 0xF7, 0xFA)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(10)

    _add_pptx_footer(slide, prs, scenario_name, page_num)


def _add_pptx_chart_slide(prs, png_bytes: bytes, title: str, scenario_name: str, page_num: int) -> None:
    from pptx.util import Inches as I, Pt, Emu
    from pptx.dml.color import RGBColor
    from PIL import Image as PILImage

    C_NAVY = RGBColor(0x01, 0x24, 0x4A)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, I(0.7))
    band.fill.solid()
    band.fill.fore_color.rgb = C_NAVY
    band.line.fill.background()
    ttl = slide.shapes.add_textbox(I(0.5), I(0.12), I(12), I(0.5))
    tp = ttl.text_frame.paragraphs[0]
    tp.text = title
    for r in tp.runs:
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = C_WHITE

    img_stream = BytesIO(png_bytes)
    img = PILImage.open(img_stream)
    img_w, img_h = img.size
    img_stream.seek(0)

    avail_w = 12.5
    avail_h = 6.0
    scale = min(avail_w / (img_w / 96.0), avail_h / (img_h / 96.0), 1.0)
    disp_w = (img_w / 96.0) * scale
    disp_h = (img_h / 96.0) * scale
    left = (13.333 - disp_w) / 2
    top = 0.9 + (6.0 - disp_h) / 2

    slide.shapes.add_picture(img_stream, I(left), I(top), I(disp_w), I(disp_h))
    _add_pptx_footer(slide, prs, scenario_name, page_num)


def _add_pptx_changelog_slides(prs, change_log: list, scenario_name: str, page_start: int) -> int:
    from pptx.util import Inches as I, Pt, Emu
    from pptx.dml.color import RGBColor

    C_NAVY = RGBColor(0x01, 0x24, 0x4A)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    ROWS_PER_SLIDE = 20
    page = page_start

    if not change_log:
        return page

    chunks = [change_log[i:i + ROWS_PER_SLIDE] for i in range(0, len(change_log), ROWS_PER_SLIDE)]
    for chunk_idx, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        band = slide.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, I(0.7))
        band.fill.solid()
        band.fill.fore_color.rgb = C_NAVY
        band.line.fill.background()
        ttl = slide.shapes.add_textbox(I(0.5), I(0.12), I(12), I(0.5))
        tp = ttl.text_frame.paragraphs[0]
        suffix = f" ({chunk_idx + 1}/{len(chunks)})" if len(chunks) > 1 else ""
        tp.text = f"Change Log{suffix}"
        for r in tp.runs:
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = C_WHITE

        headers = ["Action", "Employee", "From", "To", "Field", "Timestamp"]
        table_shape = slide.shapes.add_table(
            len(chunk) + 1, len(headers),
            I(0.4), I(1.0), I(12.5), I(min(6.0, 0.3 * (len(chunk) + 1))),
        )
        tbl = table_shape.table
        for c, h in enumerate(headers):
            cell = tbl.cell(0, c)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_NAVY
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(10)
                    r.font.color.rgb = C_WHITE

        for i, entry in enumerate(chunk, start=1):
            action = str(entry.get("action", ""))
            emp = str(entry.get("emp_id", "") or "")
            old_mgr = str(entry.get("old_mgr_id", "") or "")
            new_mgr = str(entry.get("new_mgr_id", "") or "")
            field = str(entry.get("field", "") or action)
            when = (entry.get("created_at") or entry.get("timestamp") or "")[:19].replace("T", " ")
            vals = [action, emp, old_mgr, new_mgr, field, when]
            for c, v in enumerate(vals):
                cell = tbl.cell(i, c)
                cell.text = v
                if i % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xF5, 0xF7, 0xFA)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)

        _add_pptx_footer(slide, prs, scenario_name, page)
        page += 1

    return page


def _add_pptx_footer(slide, prs, scenario_name: str, page_num: int) -> None:
    from pptx.util import Inches as I, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    line = slide.shapes.add_shape(
        1, I(0.5), I(7.05), I(12.3), Pt(1.5),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0x01, 0x24, 0x4A)
    line.line.fill.background()

    foot = slide.shapes.add_textbox(I(0.5), I(7.1), I(12.3), I(0.3))
    fp = foot.text_frame.paragraphs[0]
    fp.text = f"OrgSight 2.0  |  {scenario_name}  |  Page {page_num}"
    fp.alignment = PP_ALIGN.CENTER
    for r in fp.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x8A, 0x9A, 0xB4)


def _build_scenario_pptx(
    records: list,
    dataset: dict,
    scenario: dict,
    detail: str = "summary",
) -> bytes:
    """Build a multi-slide PPTX with title, summary, scenario comparison,
    phasing monthly timeline, org chart pages, and change log.
    *detail* controls depth:
      - ``overview``: title + summary + comparison + phasing + L1-L2 overview only
      - ``summary``: same as overview (L1-L2 overview only, no subtree slides)
      - ``full``: + recursive L1 team subtrees
    """
    from pptx.util import Emu

    emp_col = dataset["emp_col"]
    mgr_col = dataset["mgr_col"]
    jtc = dataset.get("job_title_col")
    ftc = dataset.get("fte_col")
    flc = dataset.get("flc_col")
    ctc = dataset.get("country_col")

    summary_data = db_service.get_scenario_summary(scenario["id"])
    change_log = db_service.get_change_log(scenario["id"])
    tree = get_tree_structure(records, emp_col, mgr_col)

    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333 in
    prs.slide_height = Emu(6858000)   # 7.5 in

    svg_kwargs = dict(
        emp_col=emp_col, mgr_col=mgr_col,
        job_title_col=jtc, fte_col=ftc, flc_col=flc, country_col=ctc,
    )

    _add_pptx_title_slide(prs, dataset, scenario)

    _add_pptx_summary_slide(prs, summary_data, dataset, scenario, page_num=2)

    _add_pptx_comparison_slide(prs, dataset["id"], scenario["id"], scenario["name"], page_num=3)

    _add_pptx_phasing_slide(prs, scenario["id"], scenario["name"], page_num=4)

    page = 5

    overview_svg = render_scenario_svg(
        records, **svg_kwargs,
        title=f"Organization Chart — Overview (L1–L2)",
        subtitle=f"{dataset['name']}  ·  {scenario['name']}",
        max_depth=2,
    )
    try:
        overview_png = _svg_to_png(overview_svg, max_width=3840)
        _add_pptx_chart_slide(prs, overview_png, "Organization Chart — Overview (L1–L2)", scenario["name"], page)
        page += 1
    except Exception:
        pass

    if detail == "full":
        for root_id in tree["roots"]:
            l1_kids = tree["children"].get(root_id, [])
            for kid_id in l1_kids:
                rec = tree["by_id"].get(kid_id, {})
                kid_title = str(rec.get(jtc) or rec.get("Job Title") or kid_id) if jtc else str(rec.get("Job Title") or kid_id)
                kid_hc = tree["headcount"].get(kid_id, 0)
                if kid_hc < 1:
                    continue
                subtree_svg = render_scenario_svg(
                    records, **svg_kwargs,
                    title=f"{kid_title} — Team Structure",
                    subtitle=f"{kid_hc} headcount",
                    root_id=kid_id,
                    max_depth=3,
                )
                try:
                    subtree_png = _svg_to_png(subtree_svg, max_width=3840)
                    _add_pptx_chart_slide(
                        prs, subtree_png,
                        f"{kid_title} — Team Structure",
                        scenario["name"], page,
                    )
                    page += 1
                except Exception:
                    pass

                if kid_hc > 20:
                    l2_kids = tree["children"].get(kid_id, [])
                    for gk_id in l2_kids:
                        gk_rec = tree["by_id"].get(gk_id, {})
                        gk_title = str(gk_rec.get(jtc) or gk_rec.get("Job Title") or gk_id) if jtc else str(gk_rec.get("Job Title") or gk_id)
                        gk_hc = tree["headcount"].get(gk_id, 0)
                        if gk_hc < 5:
                            continue
                        deep_svg = render_scenario_svg(
                            records, **svg_kwargs,
                            title=f"{gk_title} — Detail",
                            subtitle=f"{gk_hc} headcount",
                            root_id=gk_id,
                            max_depth=4,
                        )
                        try:
                            deep_png = _svg_to_png(deep_svg, max_width=3840)
                            _add_pptx_chart_slide(
                                prs, deep_png,
                                f"{gk_title} — Detail",
                                scenario["name"], page,
                            )
                            page += 1
                        except Exception:
                            pass

    if change_log:
        page = _add_pptx_changelog_slides(prs, change_log, scenario["name"], page)

    out = BytesIO()
    prs.save(out)
    return out.getvalue()


# ===================================================================
# Pipeline endpoints (stateless -- data in body, no DB interaction)
# ===================================================================

@router.post("/upload")
async def upload(
    file: UploadFile,
    request: Request,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        contents = await file.read()
        df = read_excel_file(BytesIO(contents))
        rows_count = len(df)
        records = df.to_dict(orient="records")
        safe_records = jsonable_encoder(records)
        write_activity_log(
            username=username, action="process", module="Upload",
            rows_output=rows_count, status="success",
            details=f"Uploaded file: {file.filename}",
        )
        return {"columns": df.columns.tolist(), "records": safe_records}
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="Upload",
            status="error", details=str(e),
        )
        raise


@router.post("/cleanup", response_class=ORJSONResponse)
def cleanup_endpoint(
    payload: List[Dict],
    project_id: int,
    remove_exclusion: bool = True,
    country_col: Optional[str] = None,
    dataset_id: Optional[int] = Query(None),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        df = pd.DataFrame(payload)
        rows_input = len(df)
        df = build_country_flag(df, country_col=country_col)
        df, removed_count = apply_exclusion_filter(df, remove=remove_exclusion)
        df = df.replace([np.inf, -np.inf], np.nan)
        rows_output = len(df)
        records = df.to_dict(orient="records")
        if dataset_id is not None:
            _require_dataset_in_project(dataset_id, project_id)
            db_service.touch_dataset_pipeline(dataset_id, "cleanup")
        write_activity_log(
            username=username, action="process", module="Cleanup",
            rows_input=rows_input, rows_output=rows_output, status="success",
            details=f"Removed {removed_count} exclusion rows, Country_Flag generated from '{country_col or 'not specified'}'",
        )
        return {"df": records, "removed": int(removed_count)}
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="Cleanup",
            rows_input=len(payload), status="error", details=str(e),
        )
        raise


# ===================================================================
# Smart upload pipeline
# ===================================================================

@router.post("/smart-upload")
async def smart_upload_endpoint(
    file: UploadFile,
    request: Request,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Smart upload: reads messy Excel, auto-detects headers, unmerges cells,
    cleans column names, fixes float IDs, and returns data + preprocessing summary.
    """
    username = user["username"]
    try:
        contents = await file.read()
        df, preprocessing = smart_read_excel(contents, filename=file.filename or "")
        rows_count = len(df)
        records = df.to_dict(orient="records")
        safe_records = jsonable_encoder(records)
        write_activity_log(
            username=username, action="process", module="SmartUpload",
            rows_output=rows_count, status="success",
            details=(
                f"File: {file.filename}, sheet: {preprocessing.get('sheet_used', '?')}, "
                f"header_row: {preprocessing.get('header_row_detected', 1)}, "
                f"merged_cells: {preprocessing.get('merged_cells_resolved', 0)}"
            ),
        )
        return {
            "columns": df.columns.tolist(),
            "records": safe_records,
            "preprocessing": preprocessing,
        }
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="SmartUpload",
            status="error", details=str(e),
        )
        raise


# ===================================================================
# Column auto-mapping
# ===================================================================

@router.post("/auto-map-columns")
def auto_map_columns_endpoint(
    body: ColumnMappingBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Auto-map uploaded columns to OrgSight target columns using alias
    matching + a single LLM call for unresolved columns."""
    username = user["username"]
    try:
        mappings, message, requires_attention, mapping_summary = auto_map_columns_with_feedback(body.columns, body.sample_rows)
        write_activity_log(
            username=username, action="process", module="ColumnMapping",
            status="success",
            details=f"Mapped {sum(1 for m in mappings.values() if m.get('source_column'))} of {len(mappings)} target columns",
        )
        return {
            "mappings": mappings,
            "message": message,
            "requires_attention": requires_attention,
            "mapping_summary": mapping_summary,
        }
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="ColumnMapping",
            status="error", details=str(e),
        )
        raise


# ===================================================================
# Rationalisation
# ===================================================================

@router.post("/rationalise")
def rationalise_endpoint(
    body: RationaliseBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Run three-layer rationalisation (exact + fuzzy + batch LLM) on
    Function, Subfunction, and Title columns.  Returns proposed mappings
    for user review — does NOT modify the data."""
    username = user["username"]
    try:
        result = rationalise(
            body.records,
            func_col=body.func_col,
            subfunc_col=body.subfunc_col,
            title_col=body.title_col,
            ignore_learned_aliases=body.ignore_learned_aliases,
        )
        summary = result.get("summary", {})
        write_activity_log(
            username=username, action="process", module="Rationalise",
            rows_input=len(body.records), status="success",
            details=(
                f"Functions: {summary.get('functions_total', 0)} "
                f"(exact={summary.get('functions_exact', 0)}, "
                f"fuzzy={summary.get('functions_fuzzy', 0)}, "
                f"ai={summary.get('functions_ai', 0)}), "
                f"Titles: {summary.get('titles_total', 0)} "
                f"(exact={summary.get('titles_exact', 0)}, "
                f"fuzzy={summary.get('titles_fuzzy', 0)}, "
                f"ai={summary.get('titles_ai', 0)})"
            ),
        )
        return result
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="Rationalise",
            rows_input=len(body.records), status="error", details=str(e),
        )
        raise


@router.post("/rationalise/apply", response_class=ORJSONResponse)
def rationalise_apply_endpoint(
    body: RationaliseApplyBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Apply user-approved rationalisation mappings to all rows.
    Adds six new columns: Rationalised Function/Subfunction/Title + Sources."""
    username = user["username"]
    try:
        updated = apply_rationalisation(
            body.records,
            func_col=body.func_col,
            subfunc_col=body.subfunc_col,
            title_col=body.title_col,
            approved_functions=body.approved_functions,
            approved_subfunctions=body.approved_subfunctions,
            approved_titles=body.approved_titles,
        )
        persist_approved_mappings(
            body.approved_functions,
            body.approved_subfunctions,
            body.approved_titles,
            enrich_learned=body.enrich_learned,
        )
        if body.dataset_id is not None:
            _require_dataset_in_project(body.dataset_id, project_id)
            db_service.touch_dataset_pipeline(body.dataset_id, "rationalise")
        write_activity_log(
            username=username, action="process", module="RationaliseApply",
            rows_input=len(body.records), rows_output=len(updated),
            status="success",
            details=f"Applied rationalisation to {len(updated)} rows",
        )
        return {"records": updated}
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="RationaliseApply",
            rows_input=len(body.records), status="error", details=str(e),
        )
        raise


@router.get("/learned-taxonomy")
def get_learned_taxonomy(
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Return all global learned mapping aliases for the Mapping Registry."""
    return {"entries": list_learned_taxonomy_entries()}


@router.patch("/learned-taxonomy")
def patch_learned_taxonomy(
    body: LearnedTaxonomyPatchBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Update, disable, or delete a learned mapping entry."""
    username = user["username"]
    try:
        result = update_learned_taxonomy_entry(
            body.entry_id,
            resolved=body.resolved,
            disabled=body.disabled,
            apply_to_matching=body.apply_to_matching,
            delete=body.delete,
        )
        write_activity_log(
            username=username, action="update", module="MappingRegistry",
            status="success",
            details=f"Updated learned mapping {body.entry_id}",
        )
        return result
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        write_activity_log(
            username=username, action="update", module="MappingRegistry",
            status="error", details=str(e),
        )
        raise


# ===================================================================
# Full pipeline (upload → cleanup → map → rationalise → validate)
# ===================================================================

@router.post("/process-upload")
async def process_upload_endpoint(
    file: UploadFile,
    request: Request,
    project_id: int,
    remove_exclusion: bool = Query(True),
    user: dict = Depends(require_project_access()),
):
    """One-shot pipeline: upload → preprocess → auto-map columns → cleanup →
    rationalise → validate.  Returns everything the frontend needs to render
    the Upload & Prepare screen in a single response."""
    username = user["username"]
    try:
        contents = await file.read()

        # 1. Smart read
        df, preprocessing = smart_read_excel(contents, filename=file.filename or "")
        columns = df.columns.tolist()
        records = df.to_dict(orient="records")

        # 2. Auto-map columns
        sample_rows = records[:10]
        col_mappings, mapping_msg, mapping_attention, _ = auto_map_columns_with_feedback(columns, sample_rows)

        # Determine mapped columns for downstream steps
        country_col = (col_mappings.get("country") or {}).get("source_column")
        func_col = (col_mappings.get("function") or {}).get("source_column")
        subfunc_col = (col_mappings.get("subfunction") or {}).get("source_column")
        title_col = (col_mappings.get("job_title") or {}).get("source_column")
        emp_col = (col_mappings.get("employee_id") or {}).get("source_column")
        mgr_col = (col_mappings.get("manager_id") or {}).get("source_column")

        # 3. Cleanup
        cleanup_removed = 0
        if country_col or "exclusion list" in df.columns:
            df = build_country_flag(df, country_col=country_col)
            df, cleanup_removed = apply_exclusion_filter(df, remove=remove_exclusion)
            df = df.replace([np.inf, -np.inf], np.nan)
            records = df.to_dict(orient="records")
            columns = df.columns.tolist()

        # 4. Rationalise (propose mappings)
        rationalisation_result = None
        if func_col or subfunc_col or title_col:
            rationalisation_result = rationalise(
                records,
                func_col=func_col,
                subfunc_col=subfunc_col,
                title_col=title_col,
            )

        # 5. Validate (if emp_col and mgr_col are mapped)
        validation_result = None
        if emp_col and mgr_col:
            try:
                raw_validation = validate_org_data(df, emp_col, mgr_col)
                validation_result = {
                    k: v for k, v in raw_validation.items() if k != "df_with_flags"
                }
                vdf = raw_validation.get("df_with_flags")
                if vdf is not None:
                    flag_cols = [c for c in vdf.columns if str(c).startswith("FLAG_")]
                    validation_result["flag_counts"] = {
                        c: int(vdf[c].fillna(0).astype(bool).sum())
                        for c in flag_cols
                    }
            except Exception as ve:
                logger.warning("Auto-validation failed: %s", ve)

        safe_records = jsonable_encoder(records)

        write_activity_log(
            username=username, action="process", module="ProcessUpload",
            rows_input=preprocessing.get("total_rows", len(records)),
            rows_output=len(records), status="success",
            details=(
                f"File: {file.filename}, "
                f"sheet: {preprocessing.get('sheet_used', '?')}, "
                f"cleanup_removed: {cleanup_removed}, "
                f"columns_mapped: {sum(1 for m in col_mappings.values() if m.get('source_column'))}"
            ),
        )

        return {
            "columns": columns,
            "records": safe_records,
            "preprocessing": preprocessing,
            "column_mappings": col_mappings,
            "column_mapping_message": mapping_msg,
            "column_mapping_requires_attention": mapping_attention,
            "cleanup": {"removed": cleanup_removed},
            "rationalisation": rationalisation_result,
            "validation": validation_result,
        }
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="ProcessUpload",
            status="error", details=str(e),
        )
        raise


@router.post("/validate")
async def validate_endpoint(
    request: Request,
    project_id: int,
    emp_col: str = Query(...),
    mgr_col: str = Query(...),
    span_col: str | None = Query(None),
    download: bool = Query(False),
    dataset_id: Optional[int] = Query(None),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        payload = await request.json()
        df = pd.DataFrame(payload)
        rows_input = len(df)
        result = validate_org_data(df, emp_col, mgr_col, span_col)

        if download:
            excel_bytes = export_excel(result["df_with_flags"], sheet_name="Data_With_Flags")
            write_activity_log(
                username=username, action="download", module="Validate",
                rows_input=rows_input, rows_output=len(result["df_with_flags"]),
                status="success", details="Downloaded validated data with flags",
            )
            return StreamingResponse(
                BytesIO(excel_bytes),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": "attachment; filename=validated_org_data.xlsx"},
            )

        df_flags = result["df_with_flags"].replace([np.inf, -np.inf], np.nan)
        records = df_flags.where(pd.notnull(df_flags), None).to_dict(orient="records")
        circular_ids = result.get("circular_reference_ids", [])

        if dataset_id is not None and not download:
            _require_dataset_in_project(dataset_id, project_id)
            db_service.touch_dataset_pipeline(dataset_id, "validate")

        write_activity_log(
            username=username, action="process", module="Validate",
            rows_input=rows_input, rows_output=len(records), status="success",
            details=(
                f"Found {len(result['duplicate_ids'])} duplicates, "
                f"{len(result['missing_manager_ids'])} missing managers, "
                f"{len(result['invalid_manager_ids'])} invalid managers, "
                f"{len(circular_ids)} circular references"
            ),
        )
        return ORJSONResponse({
            "duplicate_ids": result["duplicate_ids"],
            "missing_manager_ids": result["missing_manager_ids"],
            "invalid_manager_ids": result["invalid_manager_ids"],
            "circular_reference_ids": circular_ids,
            "top_manager": result["top_manager"],
            "df_with_flags": records,
        })
    except Exception as e:
        write_activity_log(
            username=username, action="download" if download else "process",
            module="Validate", status="error", details=str(e),
        )
        raise


@router.post("/completeness_heatmap")
async def completeness_heatmap_endpoint(
    body: CompletenessHeatmapBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Return a group × field completeness matrix for heatmap visualization."""
    username = user["username"]
    try:
        df = pd.DataFrame(body.data)
        if df.empty:
            raise HTTPException(status_code=400, detail="No data provided")
        if not body.fields:
            raise HTTPException(status_code=400, detail="Select at least one field to analyze")
        if not body.group_col:
            raise HTTPException(status_code=400, detail="Select a group-by column")

        result = get_completeness_matrix(df, body.fields, body.group_col)

        write_activity_log(
            username=username,
            action="process",
            module="Validate",
            rows_input=len(df),
            rows_output=len(result["matrix"]),
            status="success",
            details=f"Completeness heatmap: {result['summary']['fields_count']} fields × {result['summary']['groups_count']} groups",
        )
        return result
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        write_activity_log(
            username=username,
            action="process",
            module="Validate",
            status="error",
            details=str(e),
        )
        raise


@router.post("/filter_errors")
async def filter_invalids(
    request: Request,
    project_id: int,
    emp_col: str = Query(...),
    mgr_col: str = Query(...),
    remove_dup: bool = Query(True),
    remove_missing: bool = Query(True),
    remove_invalid: bool = Query(True),
    remove_circular: bool = Query(True),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        payload = await request.json()
        df = pd.DataFrame(payload)
        original_count = len(df)
        df_filtered = filter_errors(
            df, emp_col, mgr_col,
            remove_dup=remove_dup, remove_missing=remove_missing,
            remove_invalid=remove_invalid, remove_circular=remove_circular,
        )
        filtered_count = len(df_filtered)
        removed_count = original_count - filtered_count
        df_filtered = df_filtered.replace([np.inf, -np.inf], np.nan)
        records = df_filtered.where(pd.notnull(df_filtered), None).to_dict(orient="records")
        write_activity_log(
            username=username, action="process", module="Filter Errors",
            rows_input=original_count, rows_output=filtered_count,
            status="success", details=f"Removed {removed_count} error rows",
        )
        return ORJSONResponse({
            "df": records, "original_count": original_count,
            "filtered_count": filtered_count, "removed_count": removed_count,
        })
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="Filter Errors",
            status="error", details=str(e),
        )
        raise


@router.post("/hierarchy")
async def hierarchy_endpoint(
    request: Request,
    project_id: int,
    emp_col: str = Query(...),
    mgr_col: str = Query(...),
    flc_col: str | None = Query(None),
    fte_col: str | None = Query(None),
    job_title_col: str | None = Query(None),
    download: bool = Query(False),
    dataset_id: int | None = Query(None),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        payload = await request.json()
        df = pd.DataFrame(payload)
        rows_input = len(df)

        if emp_col not in df.columns or mgr_col not in df.columns:
            raise ValueError(f"Selected columns '{emp_col}' or '{mgr_col}' not in dataframe")

        df = compute_levels(df, emp_col, mgr_col)
        df = compute_direct_span(df, emp_col, mgr_col)
        df, max_depth = compute_chains(df, emp_col, mgr_col)
        df = compute_total_reports(df, emp_col, mgr_col)
        df = compute_avg_flc(df, flc_col, fte_col)

        # Apply user-defined formula columns if a saved dataset is active
        if dataset_id:
            try:
                formulas = db_service.list_formulas(dataset_id)
                if formulas:
                    enriched = apply_formulas_to_records(df.to_dict(orient="records"), formulas)
                    df = pd.DataFrame(enriched)
            except Exception:
                pass  # Formula errors must not break the hierarchy result

        df["Last_Employee"] = df["Chain_reversed"].apply(
            lambda x: x[-1] if isinstance(x, list) and len(x) > 0 else None
        )
        level_cols = [f"L{i+1}" for i in range(max_depth)] if max_depth > 0 else []

        # Compact preview columns (UI only) — system hierarchy view
        potential_base_cols = []
        if job_title_col and job_title_col in df.columns:
            potential_base_cols.append(job_title_col)
        elif "Job Title" in df.columns and not job_title_col:
            potential_base_cols.append("Job Title")
        for col in ["Division (Reporting Line)", "Country", "Level"]:
            if col in df.columns and col not in potential_base_cols:
                potential_base_cols.append(col)

        preview_cols = level_cols + ["Last_Employee"] + potential_base_cols + ["Total_Reports", "Avg_FLC"]
        preview_cols = [c for c in preview_cols if c in df.columns]
        df_preview = df[preview_cols].copy()

        from services.hierarchy_service import sort_hierarchy
        df_preview = sort_hierarchy(df_preview, max_depth)
        df_preview = df_preview.replace([pd.NA, np.nan, np.inf, -np.inf], "")

        # Full export: preserve ALL original columns, append system-generated ones
        SYSTEM_COLS = (
            ["Level", "Span", "Total_Reports", "Avg_FLC", "Last_Employee", "Chain", "Chain_reversed"]
            + level_cols
        )
        original_cols = [c for c in df.columns if c not in SYSTEM_COLS]
        # Prefer original order, then append any system cols that exist
        export_cols = original_cols + [c for c in SYSTEM_COLS if c in df.columns]
        df_export = df[export_cols].copy()
        df_export = sort_hierarchy(df_export, max_depth)
        # Excel-safe: stringify list columns (Chain / Chain_reversed)
        for list_col in ("Chain", "Chain_reversed"):
            if list_col in df_export.columns:
                df_export[list_col] = df_export[list_col].apply(
                    lambda v: " > ".join(str(x) for x in v) if isinstance(v, list) else v
                )
        df_export = df_export.replace([pd.NA, np.nan, np.inf, -np.inf], "")

        if download:
            excel_bytes = export_excel(df_export, sheet_name="Hierarchy")
            write_activity_log(
                username=username, action="download", module="Hierarchy",
                rows_input=rows_input, rows_output=len(df_export),
                status="success", details=f"Downloaded hierarchy with max depth {max_depth}",
            )
            return StreamingResponse(
                BytesIO(excel_bytes),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": "attachment; filename=org_hierarchy_output.xlsx"},
            )

        write_activity_log(
            username=username, action="process", module="Hierarchy",
            rows_input=rows_input, rows_output=len(df),
            status="success", details=f"Processed hierarchy with max depth {max_depth}",
        )
        # df keeps every original column + appended system fields for downstream
        # analysis/save; preview stays compact for the Hierarchy UI table.
        return ORJSONResponse({
            "preview": df_preview.head(20).to_dict(orient="records"),
            "df": df.replace([pd.NA, np.nan, np.inf, -np.inf], "").to_dict(orient="records"),
            "rows_processed": len(df), "max_depth": max_depth,
        })
    except Exception as e:
        write_activity_log(
            username=username, action="download" if download else "process",
            module="Hierarchy", status="error", details=str(e),
        )
        raise


# ---------------------------------------------------------------------------
# Formula CRUD endpoints (Feature 7)
# ---------------------------------------------------------------------------

@router.get("/datasets/{dataset_id}/formulas")
def list_formulas(
    dataset_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    return {"formulas": db_service.list_formulas(dataset_id)}


@router.post("/datasets/{dataset_id}/formulas")
def create_formula(
    dataset_id: int,
    body: FormulaCreateBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    # Validate the expression can be parsed (columns unknown here; just syntax check)
    validation = formula_validate_expression(body.expression, [])
    # A syntax error with no columns mentioned will have error "Unknown columns:..."
    # which is acceptable — we allow it through (columns exist at runtime).
    # Only reject hard parse errors.
    if not validation["valid"] and "Unknown column(s)" not in (validation["error"] or ""):
        raise HTTPException(status_code=400, detail=validation["error"])
    formula = db_service.create_formula(dataset_id, body.col_name, body.expression)
    return {"formula": formula}


@router.delete("/datasets/{dataset_id}/formulas/{formula_id}")
def delete_formula(
    dataset_id: int,
    formula_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    db_service.delete_formula(formula_id)
    return {"status": "deleted"}


@router.post("/datasets/{dataset_id}/formulas/preview")
def preview_formula(
    dataset_id: int,
    body: FormulaPreviewBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    sample = body.data[: body.sample_size]
    available = body.available_columns or (list(sample[0].keys()) if sample else [])
    validation = formula_validate_expression(body.expression, available)
    if not validation["valid"]:
        raise HTTPException(status_code=400, detail=validation["error"])
    results = []
    for record in sample:
        from services.formula_service import evaluate_formula
        val = evaluate_formula(body.expression, record)
        results.append({"record": record, "result": val})
    return {"preview": results, "validation": validation}


@router.post("/datasets/{dataset_id}/formulas/validate")
def validate_formula_expression(
    dataset_id: int,
    body: FormulaPreviewBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    available = body.available_columns or []
    return formula_validate_expression(body.expression, available)


@router.post("/spans_layers")
async def spans_layers_endpoint(
    request: Request,
    project_id: int,
    threshold: float = Query(0.0),
    download: bool = Query(False),
    emp_col: str | None = Query(None),
    mgr_col: str | None = Query(None),
    fte_col: str | None = Query(None),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        payload = await request.json()
        df = pd.DataFrame(payload)
        rows_input = len(df)

        if "Span" not in df.columns or "Level" not in df.columns:
            raise ValueError("Required columns 'Span' or 'Level' missing")

        summary = spans_and_layers(df)
        df_threshold = None
        high = low = None
        if threshold > 0:
            df_threshold, high, low = span_threshold(df, threshold=threshold)

        if download:
            if threshold <= 0:
                raise ValueError("Threshold must be greater than 0 for download")
            if df_threshold is None:
                raise ValueError("Threshold data not computed")
            excel_bytes = export_excel(df_threshold.fillna(""), sheet_name="Spans_and_Layers")
            write_activity_log(
                username=username, action="download", module="Spans & Layers",
                rows_input=rows_input, rows_output=len(df_threshold),
                status="success", details=f"Downloaded with threshold {threshold} (High: {high}, Low: {low})",
            )
            return StreamingResponse(
                BytesIO(excel_bytes),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": "attachment; filename=spans_layers_with_threshold.xlsx"},
            )

        df_out = df_threshold if df_threshold is not None else df
        df_out = df_out.replace([np.inf, -np.inf], np.nan)
        summary = summary.replace([np.inf, -np.inf], np.nan)

        insights = None
        try:
            insights = get_insights(
                df_out, threshold=threshold,
                emp_col=emp_col, mgr_col=mgr_col, fte_col=fte_col,
            )
        except Exception as insights_err:
            write_activity_log(
                username=username, action="process", module="Spans & Layers",
                status="warning", details=f"insights computation failed: {insights_err}",
            )

        write_activity_log(
            username=username, action="process", module="Spans & Layers",
            rows_input=rows_input, rows_output=len(df_out),
            status="success", details=f"Computed with threshold {threshold}",
        )
        return ORJSONResponse({
            "summary": summary.to_dict(orient="records"),
            "df": df_out.to_dict(orient="records"),
            "high": high, "low": low, "rows_processed": len(df_out),
            "insights": insights,
        })
    except Exception as e:
        write_activity_log(
            username=username, action="download" if download else "process",
            module="Spans & Layers", status="error", details=str(e),
        )
        raise


@router.post("/spans_layers/layer_employees")
async def spans_layers_layer_employees(
    request: Request,
    project_id: int,
    level: float = Query(...),
    emp_col: str | None = Query(None),
    user: dict = Depends(require_project_access()),
):
    payload = await request.json()
    df = pd.DataFrame(payload)
    if "Level" not in df.columns:
        raise ValueError("Level column missing")
    rows = get_layer_employees(df, level, emp_col=emp_col)
    return {"level": level, "employees": rows, "count": len(rows)}


@router.post("/spans_layers/manager_detail")
async def spans_layers_manager_detail(
    request: Request,
    project_id: int,
    emp_id: str = Query(...),
    threshold: float = Query(0.0),
    emp_col: str | None = Query(None),
    mgr_col: str | None = Query(None),
    fte_col: str | None = Query(None),
    user: dict = Depends(require_project_access()),
):
    payload = await request.json()
    df = pd.DataFrame(payload)
    detail = get_manager_detail(
        df, emp_id, threshold=threshold,
        emp_col=emp_col, mgr_col=mgr_col, fte_col=fte_col,
    )
    if detail is None:
        raise HTTPException(status_code=404, detail="Manager not found")
    return detail


@router.post("/crosstab")
def crosstab_endpoint(
    df: list[dict],
    project_id: int,
    col_x: str | None = None,
    col_y: str | None = None,
    fte_col: str | None = None,
    flc_col: str | None = None,
    download: bool = Query(False),
    col_x_threshold: float | None = Query(None),
    col_y_threshold: float | None = Query(None),
    threshold_metric: str | None = Query(None),
    excluded_categories: str | None = Query(None),
    preview_only: bool = Query(False),
    dataset_id: int | None = Query(None),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        # Apply formula columns before building the crosstab
        if dataset_id:
            try:
                formulas = db_service.list_formulas(dataset_id)
                if formulas:
                    df = apply_formulas_to_records(df, formulas)
            except Exception:
                pass
        df = pd.DataFrame(df)
        rows_input = len(df)
        col_x = col_x if col_x in df.columns else None
        col_y = col_y if col_y in df.columns else None
        fte_col = fte_col if fte_col in df.columns else None
        flc_col = flc_col if flc_col in df.columns else None

        excluded_list = []
        if excluded_categories:
            excluded_list = [cat.strip() for cat in excluded_categories.split(",") if cat.strip()]

        if preview_only:
            preview_data = generate_preview_data(
                df, col_x, col_y, fte_col, flc_col,
                threshold_metric, col_x_threshold, col_y_threshold,
            )
            return {"preview_data": preview_data}

        if (col_x_threshold is not None or col_y_threshold is not None) and threshold_metric:
            df = apply_others_grouping(
                df, col_x, col_y, fte_col, flc_col,
                col_x_threshold, col_y_threshold, threshold_metric, excluded_list,
            )

        result_df = generate_crosstab(df, col_x, col_y, fte_col, flc_col)

        if download:
            excel_bytes = export_excel(result_df, sheet_name="Crosstab")
            write_activity_log(
                username=username, action="download", module="Crosstab",
                rows_input=rows_input, rows_output=len(result_df),
                status="success", details=f"Crosstab by {col_x} x {col_y}",
            )
            return StreamingResponse(
                BytesIO(excel_bytes),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": "attachment; filename=crosstab.xlsx"},
            )

        result_df = result_df.replace([np.inf, -np.inf], np.nan)
        result_df = result_df.where(pd.notnull(result_df), None)

        if isinstance(result_df.columns, pd.MultiIndex):
            columns = [{"x": str(cat), "metric": str(metric)} for cat, metric in result_df.columns]
            is_multiindex = True
        else:
            columns = [{"name": str(c)} for c in result_df.columns]
            is_multiindex = False

        write_activity_log(
            username=username, action="process", module="Crosstab",
            rows_input=rows_input, rows_output=len(result_df),
            status="success", details=f"Generated crosstab by {col_x} x {col_y}",
        )
        return {
            "columns": columns, "index": result_df.index.tolist(),
            "crosstab": result_df.to_numpy().tolist(), "is_multiindex": is_multiindex,
        }
    except Exception as e:
        write_activity_log(
            username=username, action="download" if download else "process",
            module="Crosstab", status="error", details=str(e),
        )
        raise


@router.post("/orgchart")
def orgchart_endpoint(
    df: List[Dict],
    project_id: int,
    emp_col: str = Query(...),
    mgr_col: str = Query(...),
    filter1_col: Optional[str] = Query(None),
    filter1_values: Optional[str] = Query(None),
    filter2_col: Optional[str] = Query(None),
    filter2_values: Optional[str] = Query(None),
    filter3_col: Optional[str] = Query(None),
    filter3_values: Optional[str] = Query(None),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        df = pd.DataFrame(df)
        rows_input = len(df)
        for col in df.select_dtypes(include=[np.number]).columns:
            df[col] = df[col].replace([np.inf, -np.inf], np.nan).fillna(0)

        filtered_df = df.copy()
        filters_applied = False
        for fcol, fvals in [(filter1_col, filter1_values), (filter2_col, filter2_values), (filter3_col, filter3_values)]:
            if fcol and fvals:
                vals = [v.strip() for v in fvals.split(",") if v.strip()]
                if vals:
                    filtered_df = filtered_df[filtered_df[fcol].isin(vals)]
                    filters_applied = True

        if filters_applied and len(filtered_df) < len(df):
            df_with_ancestors = build_tree_preserve_ancestors(df, filtered_df, emp_col, mgr_col)
        else:
            df_with_ancestors = filtered_df

        tree = build_org_tree(df_with_ancestors, emp_col, mgr_col, limit_level=None)
        serializable_tree = make_json_serializable(tree)
        write_activity_log(
            username=username, action="process", module="Org Chart",
            rows_input=rows_input, rows_output=len(df_with_ancestors),
            status="success", details="Generated org chart tree",
        )
        return jsonable_encoder(serializable_tree)
    except Exception as e:
        write_activity_log(
            username=username, action="process", module="Org Chart",
            status="error", details=str(e),
        )
        raise


# ===================================================================
# Export endpoints
# ===================================================================

@router.post("/export_ppt")
async def export_ppt(
    request_body: dict,
    request: Request,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        svg_content = request_body.get("svg_content")
        if not svg_content:
            raise HTTPException(status_code=400, detail="svg_content is required")

        import re
        width_match = re.search(r'width="(\d+(?:\.\d+)?)"', svg_content)
        height_match = re.search(r'height="(\d+(?:\.\d+)?)"', svg_content)
        original_width = float(width_match.group(1)) if width_match else 1000
        original_height = float(height_match.group(1)) if height_match else 1000

        max_dimension = 5000
        scale_factor = 1.0
        if original_width > max_dimension or original_height > max_dimension:
            scale_factor = min(max_dimension / original_width, max_dimension / original_height)
            new_width = original_width * scale_factor
            new_height = original_height * scale_factor
            if width_match:
                svg_content = svg_content.replace(f'width="{original_width}"', f'width="{new_width}"')
            if height_match:
                svg_content = svg_content.replace(f'height="{original_height}"', f'height="{new_height}"')
            if "viewBox" not in svg_content:
                svg_content = svg_content.replace("<svg", f'<svg viewBox="0 0 {original_width} {original_height}"')

        with tempfile.TemporaryDirectory() as temp_dir:
            svg_file_path = os.path.join(temp_dir, "OrgChart.svg")
            emf_file_path = svg_file_path.replace(".svg", ".emf")
            with open(svg_file_path, "w", encoding="utf-8") as f:
                f.write(svg_content)

            inkscape_exe_path = r"C:\Program Files\Inkscape\bin\inkscape.exe"
            if not os.path.exists(inkscape_exe_path):
                raise HTTPException(status_code=500, detail="Inkscape not found")

            inkscape_cmd = f'"{inkscape_exe_path}" "{svg_file_path}" --export-type=emf --export-filename="{emf_file_path}"'
            result = subprocess.run(inkscape_cmd, shell=True, capture_output=True, text=True)
            if result.returncode != 0:
                raise HTTPException(status_code=500, detail=f"Inkscape failed: {result.stderr}")
            if not os.path.exists(emf_file_path):
                raise HTTPException(status_code=500, detail="EMF file was not created")

            with open(emf_file_path, "rb") as emf_file:
                emf_data = emf_file.read()

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        with tempfile.NamedTemporaryFile(delete=False, suffix=".emf") as temp_emf:
            temp_emf.write(emf_data)
            temp_emf_path = temp_emf.name

        try:
            slide_width = prs.slide_width
            slide_height = prs.slide_height
            margin = Inches(0.5)
            max_width = slide_width - (2 * margin)
            max_height = slide_height - (2 * margin)
            aspect_ratio = original_width / original_height if original_height > 0 else 1
            if max_width / aspect_ratio <= max_height:
                pic_width = max_width
                pic_height = pic_width / aspect_ratio
            else:
                pic_height = max_height
                pic_width = pic_height * aspect_ratio
            left = (slide_width - pic_width) / 2
            top = (slide_height - pic_height) / 2
            slide.shapes.add_picture(temp_emf_path, left, top, width=pic_width)

            ppt_bytes = BytesIO()
            prs.save(ppt_bytes)
            ppt_bytes.seek(0)
        finally:
            if os.path.exists(temp_emf_path):
                os.unlink(temp_emf_path)

        write_activity_log(
            username=username, action="download", module="Org Chart",
            status="success", details=f"Downloaded org chart as PPT (scaled: {scale_factor:.2f}x)",
        )
        return StreamingResponse(
            ppt_bytes,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": "attachment; filename=OrgChart.pptx"},
        )
    except HTTPException:
        raise
    except Exception as e:
        write_activity_log(
            username=username, action="download", module="Org Chart",
            status="error", details=str(e),
        )
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/export/excel")
async def export_excel_endpoint(
    payload: List[Dict],
    project_id: int,
    sheet_name: str = Query("Sheet1"),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        df = pd.DataFrame(payload)
        rows_count = len(df)
        df = df.replace([np.inf, -np.inf], np.nan).fillna("")
        excel_bytes = export_excel(df, sheet_name=sheet_name)
        write_activity_log(
            username=username, action="download", module="Export",
            rows_output=rows_count, status="success",
            details=f"Downloaded {sheet_name} as Excel",
        )
        return StreamingResponse(
            BytesIO(excel_bytes),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=org_data_export.xlsx"},
        )
    except Exception as e:
        write_activity_log(
            username=username, action="download", module="Export",
            status="error", details=str(e),
        )
        raise


# ===================================================================
# Activity Logs
# ===================================================================

@router.get("/logs/activity")
def get_logs(
    project_id: int,
    limit: Optional[int] = Query(100),
    user: dict = Depends(require_project_access()),
):
    logs = get_activity_logs(limit=limit)
    return {"logs": logs}


@router.get("/logs/stats/{username}")
def get_stats(
    username: str,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    stats = get_user_stats(username)
    return stats


# ===================================================================
# DB / Baseline / Scenario endpoints -- project-scoped
# ===================================================================

@router.post("/db/save_baseline")
def db_save_baseline(
    body: SaveBaselineBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    try:
        dataset_id = db_service.save_baseline(
            name=body.name, username=username, records=body.records,
            emp_col=body.emp_col, mgr_col=body.mgr_col,
            fte_col=body.fte_col, flc_col=body.flc_col,
            job_title_col=body.job_title_col, country_col=body.country_col,
            func_col=body.func_col, subfunc_col=body.subfunc_col,
            grade_col=body.grade_col, division_col=body.division_col,
            entity_col=body.entity_col, start_date_col=body.start_date_col,
            basic_pay_col=body.basic_pay_col, contract_type_col=body.contract_type_col,
            status_col=body.status_col,
            project_id=project_id,
        )
        scenarios = db_service.list_scenarios(dataset_id)

        # Auto-load into DuckDB for the "Ask OrgSight" analytical layer
        baseline_scenario = scenarios[0] if scenarios else None
        if baseline_scenario is not None:
            try:
                duckdb_manager.load(
                    user_id=user["id"],
                    project_id=project_id,
                    dataset_id=dataset_id,
                    scenario_id=baseline_scenario["id"],
                    records=body.records,
                    scenario_updated_at=baseline_scenario.get("updated_at"),
                )
            except Exception as duck_err:
                logger.warning("DuckDB load after save_baseline failed: %s", duck_err)

        write_activity_log(
            username=username, action="db_save_baseline", module="Org Chart",
            rows_input=len(body.records), rows_output=len(body.records),
            status="success", details=f"Saved dataset id={dataset_id} ({body.name}) in project {project_id}",
        )
        return {"dataset_id": dataset_id, "scenarios": scenarios}
    except Exception as e:
        write_activity_log(
            username=username, action="db_save_baseline", module="Org Chart",
            status="error", details=str(e),
        )
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/db/datasets")
def db_list_datasets(
    project_id: int,
    mine_only: bool = Query(False),
    include_preview: bool = Query(False),
    include_meta: bool = Query(True),
    user: dict = Depends(require_project_access()),
):
    username = user["username"]
    datasets = db_service.list_datasets(
        username=username if mine_only else None,
        project_id=project_id,
    )
    locks = dataset_lock_service.get_locks_for_project(project_id)
    dataset_ids = [ds["id"] for ds in datasets] if datasets else []

    previews = {}
    if include_preview and dataset_ids:
        datasets_by_id = {ds["id"]: ds for ds in datasets}
        previews = db_service.get_datasets_previews(dataset_ids, datasets_by_id)

    metas = {}
    if (include_meta or include_preview) and dataset_ids:
        metas = db_service.get_datasets_meta(dataset_ids)

    for ds in datasets:
        lock = locks.get(ds["id"])
        if lock:
            ds["locked_by"] = lock["username"]
            ds["locked_by_id"] = lock["user_id"]
        else:
            ds["locked_by"] = None
            ds["locked_by_id"] = None
        if include_preview:
            ds["preview"] = previews.get(
                ds["id"],
                {"root": None, "children": [], "total_children": 0}
            )
        if include_meta or include_preview:
            ds.update(metas.get(
                ds["id"],
                {
                    "scenario_count": 0,
                    "last_modified_at": None,
                    "promoted_scenario_name": None,
                }
            ))
    return {"datasets": datasets}


@router.get("/db/datasets/{dataset_id}")
def db_get_dataset(
    dataset_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    ds = _require_dataset_in_project(dataset_id, project_id)
    return {"dataset": ds, "scenarios": db_service.list_scenarios(dataset_id)}


@router.patch("/db/datasets/{dataset_id}/column-config")
def db_update_column_config(
    dataset_id: int,
    body: ColumnConfigBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    columns = body.dict(exclude_unset=True)
    db_service.update_dataset_column_config(dataset_id, columns)
    return {"status": "updated", "dataset_id": dataset_id}


@router.delete("/db/datasets/{dataset_id}")
def db_delete_dataset(
    dataset_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    db_service.delete_dataset(dataset_id)
    return {"status": "deleted"}




@router.get("/db/datasets/{dataset_id}/columns")
def db_dataset_columns(
    dataset_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    try:
        return db_service.get_dataset_columns(dataset_id)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/db/datasets/{dataset_id}/rate_cards")
def db_list_rate_cards(
    dataset_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    return {"rate_cards": db_service.list_rate_cards(dataset_id)}


@router.post("/db/datasets/{dataset_id}/rate_cards/preview")
def db_preview_rate_card(
    dataset_id: int,
    body: RateCardPreviewBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    try:
        rows = db_service.preview_rate_card(
            dataset_id,
            body.property_cols,
            cost_col=body.cost_col,
            min_sample=body.min_sample,
        )
        return {"rows": rows}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/datasets/{dataset_id}/rate_cards/generate")
def db_generate_rate_card(
    dataset_id: int,
    body: RateCardGenerateBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder_for_dataset()),
):
    _require_dataset_in_project(dataset_id, project_id)
    try:
        rate_card_id = db_service.create_rate_card_from_baseline(
            dataset_id=dataset_id,
            name=body.name,
            property_cols=body.property_cols,
            cost_col=body.cost_col,
            min_sample=body.min_sample,
            created_by=user["username"],
        )
        return {"rate_card": db_service.get_rate_card(rate_card_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/datasets/{dataset_id}/rate_cards/upload")
async def db_upload_rate_card(
    dataset_id: int,
    project_id: int,
    request: Request,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder_for_dataset()),
):
    from services.rate_card_service import parse_property_cols, rows_from_upload_df

    _require_dataset_in_project(dataset_id, project_id)
    form = await request.form()
    upload = form.get("file")
    name = str(form.get("name") or "Uploaded rate card")
    property_cols_raw = form.get("property_cols")
    cost_col = form.get("cost_col")
    if not upload:
        raise HTTPException(status_code=400, detail="file is required")
    try:
        content = await upload.read()
        df = pd.read_excel(BytesIO(content))
        rows = rows_from_upload_df(df)
        property_cols = parse_property_cols(property_cols_raw or "[]")
        ds = db_service.get_dataset(dataset_id)
        resolved_cost = str(cost_col or (ds or {}).get("flc_col") or "FLC")
        rate_card_id = db_service.create_rate_card_from_rows(
            dataset_id=dataset_id,
            name=name,
            property_cols=property_cols,
            cost_col=resolved_cost,
            rows=rows,
            created_by=user["username"],
        )
        return {"rate_card": db_service.get_rate_card(rate_card_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse upload: {e}")


@router.get("/db/rate_cards/{rate_card_id}")
def db_get_rate_card(
    rate_card_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    rc = db_service.get_rate_card(rate_card_id)
    if not rc:
        raise HTTPException(status_code=404, detail="Rate card not found")
    ds = db_service.get_dataset(rc["dataset_id"])
    if not ds or ds.get("project_id") != project_id:
        raise HTTPException(status_code=404, detail="Rate card not found in this project")
    return {"rate_card": rc}


@router.patch("/db/datasets/{dataset_id}/rate_cards/{rate_card_id}/rows")
def db_patch_rate_card_row(
    dataset_id: int,
    rate_card_id: int,
    body: RateCardRowBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder_for_dataset()),
):
    _require_dataset_in_project(dataset_id, project_id)
    rc = db_service.get_rate_card(rate_card_id)
    if not rc or rc["dataset_id"] != dataset_id:
        raise HTTPException(status_code=404, detail="Rate card not found")
    try:
        db_service.upsert_rate_card_row(
            rate_card_id,
            body.composite_key,
            p25=body.p25,
            p50=body.p50,
            p75=body.p75,
            avg_cost=body.avg_cost,
        )
        return {"rate_card": db_service.get_rate_card(rate_card_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.patch("/db/scenarios/{scenario_id}/rate_card")
def db_set_scenario_rate_card(
    scenario_id: int,
    body: ScenarioRateCardBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    try:
        scenario = db_service.set_scenario_rate_card(
            scenario_id,
            body.rate_card_id,
            body.rate_card_quartile,
        )
        return {"scenario": scenario}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/scenarios/{scenario_id}/rate_card/lookup")
def db_lookup_scenario_rate_card(
    scenario_id: int,
    body: RateCardLookupBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_scenario_in_project(scenario_id, project_id)
    result = db_service.lookup_rate_card_cost(scenario_id, body.values)
    return {"lookup": result}

@router.get("/db/datasets/{dataset_id}/baseline")
def db_get_baseline(
    dataset_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    return {"records": db_service.get_baseline_records(dataset_id)}


@router.get("/db/datasets/{dataset_id}/records")
def db_get_dataset_flat_records(
    dataset_id: int,
    project_id: int,
    scenario_id: Optional[int] = Query(None),
    user: dict = Depends(require_project_access()),
):
    """Return a flat record array + column list for the analytics pipeline.

    Optionally pass ?scenario_id=<id> to get the scenario's active merged
    state instead of the baseline.
    """
    _require_dataset_in_project(dataset_id, project_id)
    if scenario_id is not None:
        scenario = db_service.get_scenario(scenario_id)
        if not scenario or scenario.get("dataset_id") != dataset_id:
            raise HTTPException(status_code=404, detail="Scenario not found in dataset")
    result = db_service.get_dataset_flat_records(dataset_id, scenario_id)

    # Auto-load into DuckDB when a saved dataset is activated
    target_scenario_id = scenario_id
    if target_scenario_id is None:
        scenarios = db_service.list_scenarios(dataset_id)
        if scenarios:
            target_scenario_id = scenarios[0]["id"]
    if target_scenario_id is not None:
        try:
            sc_meta = db_service.get_scenario(target_scenario_id) or {}
            duckdb_manager.load(
                user_id=user["id"],
                project_id=project_id,
                dataset_id=dataset_id,
                scenario_id=target_scenario_id,
                records=result.get("records", []),
                scenario_updated_at=sc_meta.get("updated_at"),
            )
        except Exception as duck_err:
            logger.warning("DuckDB load on dataset fetch failed: %s", duck_err)

    return result


@router.get("/db/datasets/{dataset_id}/recent-changes")
def db_get_dataset_recent_changes(
    dataset_id: int,
    project_id: int,
    since: Optional[str] = Query(None),
    limit: int = Query(50, ge=1, le=500),
    user: dict = Depends(require_project_access()),
):
    """Activity feed for a dataset -- aggregates change_log across all
    scenarios and flags entries that are unseen by the current user."""
    _require_dataset_in_project(dataset_id, project_id)
    return db_service.get_dataset_recent_changes(
        dataset_id,
        user["id"],
        since=since,
        limit=limit,
        current_username=user["username"],
    )


@router.post("/db/datasets/{dataset_id}/mark-seen")
def db_mark_dataset_seen(
    dataset_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    ts = db_service.mark_dataset_seen(dataset_id, user["id"])
    return {"status": "ok", "last_seen": ts}


# ---------------------------------------------------------------------------
# Dataset locks -- edit-mode concurrency control
# ---------------------------------------------------------------------------

@router.post("/db/datasets/{dataset_id}/lock")
def db_dataset_acquire_lock(
    dataset_id: int,
    project_id: int,
    request: Request,
    user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    result = dataset_lock_service.acquire_lock(dataset_id, project_id, user["id"], user["username"])
    if result["acquired"]:
        write_audit_log(
            user_id=user["id"],
            action="dataset_lock_acquired",
            resource_type="dataset",
            resource_id=dataset_id,
            details={"project_id": project_id},
            ip_address=request.client.host if request.client else None,
        )
    return result


@router.post("/db/datasets/{dataset_id}/lock/heartbeat")
def db_dataset_lock_heartbeat(
    dataset_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Refresh heartbeat. No audit logging -- high frequency."""
    result = dataset_lock_service.heartbeat(
        dataset_id, user["id"], username=user["username"], project_id=project_id,
    )
    if result is None:
        raise HTTPException(404, detail="No active lock held by you for this dataset")
    return result


@router.delete("/db/datasets/{dataset_id}/lock")
def db_dataset_release_lock(
    dataset_id: int,
    project_id: int,
    request: Request,
    user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    released = dataset_lock_service.release_lock(dataset_id, user["id"])
    if released:
        write_audit_log(
            user_id=user["id"],
            action="dataset_lock_released",
            resource_type="dataset",
            resource_id=dataset_id,
            details={"project_id": project_id},
            ip_address=request.client.host if request.client else None,
        )
    return {"status": "released" if released else "no_lock"}


@router.get("/db/datasets/{dataset_id}/lock")
def db_dataset_lock_status(
    dataset_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    lock = dataset_lock_service.get_lock(dataset_id)
    if not lock:
        return {"locked": False}
    return {
        "locked": True,
        "holder": lock["username"],
        "holder_id": lock["user_id"],
        "acquired_at": lock["acquired_at"],
        "last_heartbeat": lock["last_heartbeat"],
        "is_mine": lock["user_id"] == user["id"],
    }


# ---------------------------------------------------------------------------
# Scenario CRUD -- create is gated by dataset lock
# ---------------------------------------------------------------------------

@router.post("/db/datasets/{dataset_id}/scenarios")
def db_create_scenario(
    dataset_id: int,
    body: ScenarioBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder_for_dataset()),
):
    _require_dataset_in_project(dataset_id, project_id)
    username = user["username"]
    try:
        scenario_id = db_service.create_scenario(
            dataset_id=dataset_id, name=body.name,
            description=body.description or "",
            source_scenario_id=body.source_scenario_id,
            rate_card_id=body.rate_card_id,
            rate_card_quartile=body.rate_card_quartile or "p50",
        )
        write_activity_log(
            username=username, action="scenario_create", module="Org Chart",
            status="success", details=f"Created scenario {scenario_id} ({body.name})",
        )
        return {"scenario": db_service.get_scenario(scenario_id)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/db/datasets/{dataset_id}/compare")
def db_compare_scenarios(
    dataset_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_dataset_in_project(dataset_id, project_id)
    scenarios = db_service.list_scenarios(dataset_id)
    return {"scenarios": [db_service.get_scenario_summary(s["id"]) for s in scenarios]}


# ---------------------------------------------------------------------------
# Scenario endpoints -- all verify scenario belongs to project via dataset
# ---------------------------------------------------------------------------

@router.get("/db/scenarios/{scenario_id}")
def db_get_scenario(
    scenario_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    scenario, dataset = _require_scenario_in_project(scenario_id, project_id)
    records = db_service.get_scenario_records(scenario_id)

    try:
        duckdb_manager.load(
            user_id=user["id"],
            project_id=project_id,
            dataset_id=dataset["id"],
            scenario_id=scenario_id,
            records=records,
            scenario_updated_at=scenario.get("updated_at"),
        )
    except Exception as duck_err:
        logger.warning("DuckDB load on scenario fetch failed: %s", duck_err)

    return {
        "scenario": scenario,
        "records": records,
        "summary": db_service.get_scenario_summary(scenario_id),
    }


@router.patch("/db/scenarios/{scenario_id}/rename")
def db_rename_scenario(
    scenario_id: int,
    body: ScenarioBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    db_service.rename_scenario(scenario_id, body.name, body.description)
    return {"scenario": db_service.get_scenario(scenario_id)}


@router.delete("/db/scenarios/{scenario_id}")
def db_delete_scenario(
    scenario_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    db_service.delete_scenario(scenario_id)
    return {"status": "deleted"}


@router.post("/db/scenarios/{scenario_id}/move")
def db_scenario_move(
    scenario_id: int,
    body: MoveBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    username = user["username"]
    try:
        record = db_service.move_employee(
            scenario_id=scenario_id, emp_id=body.emp_id,
            new_mgr_id=body.new_mgr_id, username=username,
            effective_date=body.effective_date,
        )
        return {"record": record, "summary": db_service.get_scenario_summary(scenario_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/scenarios/{scenario_id}/edit")
def db_scenario_edit(
    scenario_id: int,
    body: EditBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    username = user["username"]
    try:
        record = db_service.edit_employee(
            scenario_id=scenario_id, emp_id=body.emp_id,
            updates=body.updates, username=username,
            effective_date=body.effective_date,
        )
        return {"record": record, "summary": db_service.get_scenario_summary(scenario_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/scenarios/{scenario_id}/add")
def db_scenario_add(
    scenario_id: int,
    body: AddBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    username = user["username"]
    try:
        record = db_service.add_employee(
            scenario_id=scenario_id, record=body.record,
            emp_id=body.emp_id, mgr_id=body.mgr_id, level=body.level,
            fte=body.fte, flc=body.flc, username=username,
            effective_date=body.effective_date,
        )
        return {"record": record, "summary": db_service.get_scenario_summary(scenario_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/scenarios/{scenario_id}/clone")
def db_scenario_clone(
    scenario_id: int,
    body: CloneBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    username = user["username"]
    try:
        record = db_service.clone_employee(
            scenario_id=scenario_id,
            source_emp_id=body.source_emp_id,
            new_emp_id=body.new_emp_id,
            new_mgr_id=body.new_mgr_id,
            username=username,
            effective_date=body.effective_date,
        )
        return {"record": record, "summary": db_service.get_scenario_summary(scenario_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/scenarios/{scenario_id}/flag")
def db_scenario_flag(
    scenario_id: int,
    body: FlagBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    username = user["username"]
    try:
        record = db_service.flag_employee(
            scenario_id=scenario_id, emp_id=body.emp_id,
            flagged=body.flagged, username=username,
            effective_date=body.effective_date,
        )
        return {"record": record, "summary": db_service.get_scenario_summary(scenario_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.patch("/db/scenarios/{scenario_id}/change_log/bulk_date")
def db_scenario_bulk_date(
    scenario_id: int,
    body: BulkDateBody,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    """Bulk-assign (or clear) an effective_date on a list of change_log rows."""
    _require_scenario_in_project(scenario_id, project_id)
    count = db_service.bulk_set_effective_date(
        scenario_id=scenario_id,
        change_ids=body.change_ids,
        effective_date=body.effective_date,
    )
    return {"updated": count}


@router.get("/db/scenarios/{scenario_id}/phasing")
def db_scenario_phasing(
    scenario_id: int,
    project_id: int,
    fy_start_month: int = Query(default=1, ge=1, le=12),
    user: dict = Depends(require_project_access()),
):
    """Return monthly phasing view for adds/removals with full-year and in-year savings."""
    _require_scenario_in_project(scenario_id, project_id)
    return db_service.get_phasing_view(scenario_id, fy_start_month=fy_start_month)


@router.get("/db/scenarios/{scenario_id}/validate")
def db_scenario_validate(
    scenario_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
):
    """Run post-hoc validation checks on the scenario's to-be state."""
    _require_scenario_in_project(scenario_id, project_id)
    issues = validate_scenario(scenario_id, project_id)
    return {"issues": issues, "count": len(issues)}


@router.post("/db/scenarios/{scenario_id}/bulk_flag")
def db_scenario_bulk_flag(
    scenario_id: int,
    project_id: int,
    body: BulkFlagBody,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    """Flag or restore a batch of employees in one transaction."""
    _require_scenario_in_project(scenario_id, project_id)
    affected = db_service.bulk_flag_employees(
        scenario_id,
        body.emp_ids,
        body.flagged,
        username=user["username"],
        effective_date=body.effective_date,
    )
    return {"affected": affected}


@router.post("/db/scenarios/{scenario_id}/bulk_edit_property")
def db_scenario_bulk_edit_property(
    scenario_id: int,
    project_id: int,
    body: BulkEditPropertyBody,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    """Set a single field to a value across multiple employees."""
    _require_scenario_in_project(scenario_id, project_id)
    affected = db_service.bulk_edit_property(
        scenario_id,
        body.emp_ids,
        body.field,
        body.value,
        username=user["username"],
        effective_date=body.effective_date,
    )
    return {"affected": affected}


@router.post("/db/scenarios/{scenario_id}/bulk_move")
def db_scenario_bulk_move(
    scenario_id: int,
    project_id: int,
    body: BulkMoveBody,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    """Reassign multiple employees to a new manager in one transaction."""
    _require_scenario_in_project(scenario_id, project_id)
    affected = db_service.bulk_move_employees(
        scenario_id,
        body.emp_ids,
        body.new_mgr_id,
        username=user["username"],
        effective_date=body.effective_date,
    )
    summary = db_service.get_scenario_summary(scenario_id)
    return {"affected": affected, "summary": summary}


@router.post("/db/scenarios/{scenario_id}/promote")
def db_scenario_promote(
    scenario_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    username = user["username"]
    try:
        db_service.promote_scenario(scenario_id)
        write_activity_log(
            username=username, action="scenario_promote", module="Org Chart",
            status="success", details=f"Promoted scenario {scenario_id} to baseline",
        )
        return {"status": "promoted", "scenario": db_service.get_scenario(scenario_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/db/scenarios/{scenario_id}/summary")
def db_scenario_summary(
    scenario_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_scenario_in_project(scenario_id, project_id)
    try:
        return db_service.get_scenario_summary(scenario_id)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.get("/db/scenarios/{scenario_id}/summary_by_dim")
def db_scenario_summary_by_dim(
    scenario_id: int,
    dim_col: str,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_scenario_in_project(scenario_id, project_id)
    try:
        breakdown = db_service.get_scenario_summary_by_dimension(scenario_id, dim_col)
        return {"breakdown": breakdown}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/scenarios/{scenario_id}/reset")
def db_scenario_reset(
    scenario_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    username = user["username"]
    try:
        db_service.reset_scenario_to_baseline(scenario_id, username=username)
        write_activity_log(
            username=username, action="scenario_reset", module="Org Chart",
            status="success", details=f"Reset scenario {scenario_id} to baseline",
        )
        return {"status": "reset", "summary": db_service.get_scenario_summary(scenario_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/db/scenarios/{scenario_id}/undo")
def db_scenario_undo(
    scenario_id: int,
    project_id: int,
    user: dict = Depends(require_project_access()),
    _lock: dict = Depends(require_dataset_lock_holder()),
):
    _require_scenario_in_project(scenario_id, project_id)
    username = user["username"]
    try:
        entry = db_service.undo_last_change(scenario_id, username=username)
        if entry is None:
            return {"status": "nothing_to_undo"}
        return {"status": "undone", "reversed": entry, "summary": db_service.get_scenario_summary(scenario_id)}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/db/scenarios/{scenario_id}/change_log")
def db_scenario_change_log(
    scenario_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    _require_scenario_in_project(scenario_id, project_id)
    return {"changes": db_service.get_change_log(scenario_id)}


# ---------------------------------------------------------------------------
# Scenario exports
# ---------------------------------------------------------------------------

@router.get("/db/scenarios/{scenario_id}/export/changes")
def db_export_change_summary(
    scenario_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    scenario, dataset = _require_scenario_in_project(scenario_id, project_id)
    summary = db_service.get_scenario_summary(scenario_id)
    changes = db_service.get_change_log(scenario_id)

    summary_rows = [
        {"Metric": "Headcount", "Baseline": summary["baseline"]["headcount"],
         "Current": summary["current"]["headcount"], "Delta": summary["delta"]["headcount"]},
        {"Metric": "Total FTE", "Baseline": round(summary["baseline"]["total_fte"], 2),
         "Current": round(summary["current"]["total_fte"], 2), "Delta": round(summary["delta"]["fte"], 2)},
        {"Metric": "Total Cost", "Baseline": round(summary["baseline"]["total_cost"], 2),
         "Current": round(summary["current"]["total_cost"], 2), "Delta": round(summary["delta"]["cost"], 2)},
        {"Metric": "Flagged-removed roles", "Baseline": 0,
         "Current": summary["flagged_removed"]["count"], "Delta": summary["flagged_removed"]["count"]},
        {"Metric": "Flagged-removed cost", "Baseline": 0,
         "Current": round(summary["flagged_removed"]["cost"], 2), "Delta": round(summary["flagged_removed"]["cost"], 2)},
        {"Metric": "Total changes logged", "Baseline": "\u2014",
         "Current": summary["change_count"], "Delta": "\u2014"},
    ]
    summary_df = pd.DataFrame(summary_rows)
    changes_df = pd.DataFrame(changes) if changes else pd.DataFrame(
        columns=["id", "scenario_id", "action", "emp_id", "old_mgr_id",
                 "new_mgr_id", "field", "old_value", "new_value", "timestamp", "username"]
    )

    def _breakdown_df(dim_col: str) -> pd.DataFrame:
        data = db_service.get_scenario_summary_by_dimension(scenario_id, dim_col)
        rows = data.get("rows", [])
        totals = data.get("totals")
        if totals:
            rows = rows + [totals]
        return pd.DataFrame([
            {
                dim_col: r["dimension_value"],
                "Baseline HC": r["baseline_hc"],
                "Baseline FTE": r["baseline_fte"],
                "Baseline Cost": r["baseline_cost"],
                "To-Be HC": r["tobe_hc"],
                "To-Be FTE": r["tobe_fte"],
                "To-Be Cost": r["tobe_cost"],
                "Delta HC": r["delta_hc"],
                "Delta FTE": r["delta_fte"],
                "Delta Cost": r["delta_cost"],
            }
            for r in rows
        ])

    def _safe_sheet_name(name: str, used: set) -> str:
        cleaned = "".join(c if c.isalnum() or c in " -_" else " " for c in name).strip()
        cleaned = cleaned[:28] or "Sheet"
        base = cleaned
        i = 1
        while cleaned.lower() in used:
            suffix = f" {i}"
            cleaned = base[: 28 - len(suffix)] + suffix
            i += 1
        used.add(cleaned.lower())
        return cleaned

    # Phasing Data
    phasing = db_service.get_phasing_view(scenario_id, fy_start_month=1)
    phasing_rows = []
    for b in phasing.get("monthly", []):
        phasing_rows.append({
            "Month": b["label"],
            "Adds": b["adds_count"],
            "Removals": b["removes_count"],
            "Edits": b["edits_count"],
            "Net HC Delta": b["net_hc_delta"],
            "Cost Delta": round(b["cost_delta"], 2),
            "Cumulative HC Delta": b["cumulative_hc_delta"],
            "Cumulative Cost Delta": round(b["cumulative_cost_delta"], 2)
        })
    phasing_df = pd.DataFrame(phasing_rows) if phasing_rows else pd.DataFrame(columns=[
        "Month", "Adds", "Removals", "Edits", "Net HC Delta", "Cost Delta", "Cumulative HC Delta", "Cumulative Cost Delta"
    ])

    # Scenario Comparison Data
    scenarios_list = db_service.list_scenarios(dataset["id"])
    comparison_rows = []
    for s in scenarios_list:
        try:
            s_sum = db_service.get_scenario_summary(s["id"])
            comparison_rows.append({
                "Scenario Name": s["name"],
                "Is Promoted": "Yes" if s["is_promoted"] else "No",
                "Baseline Headcount": s_sum["baseline"]["headcount"],
                "Current Headcount": s_sum["current"]["headcount"],
                "Delta Headcount": s_sum["delta"]["headcount"],
                "Baseline FTE": round(s_sum["baseline"]["total_fte"], 2),
                "Current FTE": round(s_sum["current"]["total_fte"], 2),
                "Delta FTE": round(s_sum["delta"]["fte"], 2),
                "Baseline Cost": round(s_sum["baseline"]["total_cost"], 2),
                "Current Cost": round(s_sum["current"]["total_cost"], 2),
                "Delta Cost": round(s_sum["delta"]["cost"], 2),
                "Changes Logged": s_sum["change_count"]
            })
        except Exception:
            continue
    comparison_df = pd.DataFrame(comparison_rows) if comparison_rows else pd.DataFrame()

    dim_cols = db_service.get_export_dimension_columns(dataset["id"])
    used_sheet_names = {"summary", "changes", "phasing", "scenario comparison"}

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        changes_df.to_excel(writer, index=False, sheet_name="Changes")
        phasing_df.to_excel(writer, index=False, sheet_name="Phasing")
        if not comparison_df.empty:
            comparison_df.to_excel(writer, index=False, sheet_name="Scenario Comparison")
        for dim_col in dim_cols:
            try:
                df = _breakdown_df(dim_col)
                sheet = _safe_sheet_name(f"By {dim_col}", used_sheet_names)
                df.to_excel(writer, index=False, sheet_name=sheet)
            except ValueError:
                continue
    buf.seek(0)

    safe_name = "".join(c for c in scenario["name"] if c.isalnum() or c in "-_") or "scenario"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="orgsight_changes_{safe_name}.xlsx"'},
    )


@router.get("/db/scenarios/{scenario_id}/export/svg")
def db_export_scenario_svg(
    scenario_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    scenario, dataset = _require_scenario_in_project(scenario_id, project_id)
    records = db_service.get_scenario_records(scenario_id)
    svg = render_scenario_svg(
        records, emp_col=dataset["emp_col"], mgr_col=dataset["mgr_col"],
        job_title_col=dataset.get("job_title_col"), fte_col=dataset.get("fte_col"),
        flc_col=dataset.get("flc_col"), country_col=dataset.get("country_col"),
        title=f"OrgSight 2.0  -  {dataset['name']}  -  {scenario['name']}",
    )
    safe_name = "".join(c for c in scenario["name"] if c.isalnum() or c in "-_") or "scenario"
    return StreamingResponse(
        BytesIO(svg.encode("utf-8")), media_type="image/svg+xml",
        headers={"Content-Disposition": f'attachment; filename="orgsight_{safe_name}.svg"'},
    )


@router.get("/db/scenarios/{scenario_id}/export/ppt")
def db_export_scenario_ppt(
    scenario_id: int,
    project_id: int,
    detail: str = Query("summary", regex="^(overview|summary|full)$"),
    _user: dict = Depends(require_project_access()),
):
    scenario, dataset = _require_scenario_in_project(scenario_id, project_id)
    records = db_service.get_scenario_records(scenario_id)

    pptx_bytes = _build_scenario_pptx(records, dataset, scenario, detail=detail)

    safe_name = "".join(c for c in scenario["name"] if c.isalnum() or c in "-_") or "scenario"
    return StreamingResponse(
        BytesIO(pptx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="orgsight_{safe_name}.pptx"'},
    )


def _render_summary_table_svg(summary, dataset, scenario) -> str:
    from html import escape
    cur = summary.get("current", {})
    base = summary.get("baseline", {})
    delta = summary.get("delta", {})
    
    phasing = db_service.get_phasing_view(scenario["id"], fy_start_month=1)
    ph_sum = phasing.get("summary", {})
    
    hc_color = '#27ae60' if delta.get('headcount', 0) <= 0 else '#c0392b'
    fte_color = '#27ae60' if delta.get('fte', 0) <= 0 else '#c0392b'
    cost_color = '#27ae60' if delta.get('cost', 0) <= 0 else '#c0392b'
    
    svg = f"""<svg xmlns="http://www.w3.org/2000/svg" width="1219" height="685" viewBox="0 0 1219 685" font-family="Inter, Arial, sans-serif">
      <rect x="0" y="0" width="1219" height="685" fill="#f4f6f9"/>
      <rect x="0" y="0" width="1219" height="60" fill="#01244a"/>
      <text x="30" y="35" font-size="20" font-weight="700" fill="#ffffff" dominant-baseline="central">Executive Summary — {escape(scenario['name'])}</text>
      
      <!-- Comparison Table -->
      <g transform="translate(50, 100)">
        <rect x="0" y="0" width="600" height="220" fill="#ffffff" rx="8" ry="8" stroke="#dce4ee" stroke-width="1.5"/>
        <rect x="0" y="0" width="600" height="40" fill="#01244a" rx="8" ry="8" clip-path="url(#table-header-clip)"/>
        <clipPath id="table-header-clip">
          <rect x="0" y="0" width="600" height="40" rx="8" ry="8"/>
        </clipPath>
        
        <text x="20" y="20" font-size="12" font-weight="700" fill="#ffffff" dominant-baseline="central">Metric</text>
        <text x="220" y="20" font-size="12" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Baseline</text>
        <text x="380" y="20" font-size="12" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Current (To-Be)</text>
        <text x="520" y="20" font-size="12" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Delta</text>
        
        <!-- Row 1: Headcount -->
        <line x1="0" y1="90" x2="600" y2="90" stroke="#dce4ee" stroke-width="1"/>
        <text x="20" y="65" font-size="11" font-weight="600" fill="#01244a" dominant-baseline="central">Headcount</text>
        <text x="220" y="65" font-size="11" fill="#4f6077" dominant-baseline="central" text-anchor="middle">{base.get('headcount', 0):,}</text>
        <text x="380" y="65" font-size="11" fill="#4f6077" dominant-baseline="central" text-anchor="middle">{cur.get('headcount', 0):,}</text>
        <text x="520" y="65" font-size="11" font-weight="700" fill="{hc_color}" dominant-baseline="central" text-anchor="middle">{delta.get('headcount', 0):+d}</text>
        
        <!-- Row 2: FTE -->
        <line x1="0" y1="140" x2="600" y2="140" stroke="#dce4ee" stroke-width="1"/>
        <text x="20" y="115" font-size="11" font-weight="600" fill="#01244a" dominant-baseline="central">Total FTE</text>
        <text x="220" y="115" font-size="11" fill="#4f6077" dominant-baseline="central" text-anchor="middle">{base.get('total_fte', 0):,.1f}</text>
        <text x="380" y="115" font-size="11" fill="#4f6077" dominant-baseline="central" text-anchor="middle">{cur.get('total_fte', 0):,.1f}</text>
        <text x="520" y="115" font-size="11" font-weight="700" fill="{fte_color}" dominant-baseline="central" text-anchor="middle">{delta.get('fte', 0):+.1f}</text>
        
        <!-- Row 3: Cost -->
        <text x="20" y="175" font-size="11" font-weight="600" fill="#01244a" dominant-baseline="central">Total Cost</text>
        <text x="220" y="175" font-size="11" fill="#4f6077" dominant-baseline="central" text-anchor="middle">${base.get('total_cost', 0):,.0f}</text>
        <text x="380" y="175" font-size="11" fill="#4f6077" dominant-baseline="central" text-anchor="middle">${cur.get('total_cost', 0):,.0f}</text>
        <text x="520" y="175" font-size="11" font-weight="700" fill="{cost_color}" dominant-baseline="central" text-anchor="middle">${delta.get('cost', 0):+,.0f}</text>
      </g>
      
      <!-- Savings & Impact Summary Card -->
      <g transform="translate(680, 100)">
        <rect x="0" y="0" width="490" height="480" fill="#fafbfc" rx="8" ry="8" stroke="#dce4ee" stroke-width="1.5"/>
        <rect x="0" y="0" width="490" height="40" fill="#01244a" rx="8" ry="8" clip-path="url(#savings-header-clip)"/>
        <clipPath id="savings-header-clip">
          <rect x="0" y="0" width="490" height="40" rx="8" ry="8"/>
        </clipPath>
        <text x="245" y="20" font-size="12" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">SAVINGS &amp; IMPACT SUMMARY</text>
        
        <text x="30" y="80" font-size="12" font-weight="700" fill="#01244a" dominant-baseline="central">Full-Year Annualized Savings:</text>
        <text x="460" y="80" font-size="12" font-weight="700" fill="#27ae60" dominant-baseline="central" text-anchor="end">${ph_sum.get('full_year_savings', 0):,.0f}</text>
        
        <text x="30" y="130" font-size="12" font-weight="700" fill="#01244a" dominant-baseline="central">Prorated In-Year Savings:</text>
        <text x="460" y="130" font-size="12" font-weight="700" fill="#27ae60" dominant-baseline="central" text-anchor="end">${ph_sum.get('in_year_savings', 0):,.0f}</text>
        
        <line x1="20" y1="175" x2="470" y2="175" stroke="#dce4ee" stroke-width="1"/>
        
        <text x="30" y="210" font-size="11" font-weight="600" fill="#4f6077" dominant-baseline="central">Months Remaining in FY:</text>
        <text x="460" y="210" font-size="11" font-weight="700" fill="#01244a" dominant-baseline="central" text-anchor="end">{ph_sum.get('months_remaining_in_fy', 0)}</text>
        
        <text x="30" y="260" font-size="11" font-weight="600" fill="#4f6077" dominant-baseline="central">Flagged Removed Roles:</text>
        <text x="460" y="260" font-size="11" font-weight="700" fill="#01244a" dominant-baseline="central" text-anchor="end">{summary.get('flagged_removed', {}).get('count', 0)}</text>
        
        <text x="30" y="310" font-size="11" font-weight="600" fill="#4f6077" dominant-baseline="central">Flagged Removed Value:</text>
        <text x="460" y="310" font-size="11" font-weight="700" fill="#01244a" dominant-baseline="central" text-anchor="end">${summary.get('flagged_removed', {}).get('cost', 0):,.0f}</text>
        
        <text x="30" y="360" font-size="11" font-weight="600" fill="#4f6077" dominant-baseline="central">Total Changes Logged:</text>
        <text x="460" y="360" font-size="11" font-weight="700" fill="#01244a" dominant-baseline="central" text-anchor="end">{summary.get('change_count', 0)}</text>
      </g>
    </svg>"""
    return svg


def _render_comparison_table_svg(dataset_id: int, active_scenario_id: int, scenario_name: str) -> str:
    from html import escape
    scenarios_list = db_service.list_scenarios(dataset_id)
    headers = ["Scenario Name", "HC", "FTE", "Cost", "Δ HC", "Δ Cost", "Changes"]
    
    rows_svg = []
    y_cursor = 100
    
    header_svg = f"""
      <rect x="50" y="{y_cursor}" width="1119" height="40" fill="#01244a" rx="4" ry="4"/>
      <text x="70" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central">Scenario Name</text>
      <text x="350" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">HC</text>
      <text x="480" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">FTE</text>
      <text x="630" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Cost</text>
      <text x="780" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Δ HC</text>
      <text x="930" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Δ Cost</text>
      <text x="1080" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Changes</text>
    """
    rows_svg.append(header_svg)
    y_cursor += 40
    
    for i, s in enumerate(scenarios_list):
        try:
            s_sum = db_service.get_scenario_summary(s["id"])
            hc = f"{s_sum['current']['headcount']:,}"
            fte = f"{s_sum['current']['total_fte']:,.1f}"
            cost = f"${s_sum['current']['total_cost']:,.0f}"
            dhc = f"{s_sum['delta']['headcount']:+d}"
            dcost = f"${s_sum['delta']['cost']:+,.0f}"
            ch = f"{s_sum['change_count']}"
            name = s["name"]
            if s["id"] == active_scenario_id:
                name += " (Active)"
            if s["is_promoted"]:
                name += " [Baseline]"
                
            bg_color = "#ffffff" if i % 2 == 0 else "#f8f9fa"
            bold_flag = 'font-weight="700" fill="#01244a"' if s["id"] == active_scenario_id else 'fill="#4f6077"'
            
            row_content = f"""
              <rect x="50" y="{y_cursor}" width="1119" height="35" fill="{bg_color}" stroke="#dce4ee" stroke-width="0.5"/>
              <text x="70" y="{y_cursor + 17.5}" font-size="10" {bold_flag} dominant-baseline="central">{escape(name)}</text>
              <text x="350" y="{y_cursor + 17.5}" font-size="10" {bold_flag} dominant-baseline="central" text-anchor="middle">{hc}</text>
              <text x="480" y="{y_cursor + 17.5}" font-size="10" {bold_flag} dominant-baseline="central" text-anchor="middle">{fte}</text>
              <text x="630" y="{y_cursor + 17.5}" font-size="10" {bold_flag} dominant-baseline="central" text-anchor="middle">{cost}</text>
              <text x="780" y="{y_cursor + 17.5}" font-size="10" {bold_flag} dominant-baseline="central" text-anchor="middle">{dhc}</text>
              <text x="930" y="{y_cursor + 17.5}" font-size="10" {bold_flag} dominant-baseline="central" text-anchor="middle">{dcost}</text>
              <text x="1080" y="{y_cursor + 17.5}" font-size="10" {bold_flag} dominant-baseline="central" text-anchor="middle">{ch}</text>
            """
            rows_svg.append(row_content)
            y_cursor += 35
        except Exception:
            continue
            
    svg = f"""<svg xmlns="http://www.w3.org/2000/svg" width="1219" height="685" viewBox="0 0 1219 685" font-family="Inter, Arial, sans-serif">
      <rect x="0" y="0" width="1219" height="685" fill="#f4f6f9"/>
      <rect x="0" y="0" width="1219" height="60" fill="#01244a"/>
      <text x="30" y="35" font-size="20" font-weight="700" fill="#ffffff" dominant-baseline="central">Scenario Comparison Analysis</text>
      {"".join(rows_svg)}
    </svg>"""
    return svg


def _render_phasing_table_svg(scenario_id: int, scenario_name: str) -> str:
    from html import escape
    phasing = db_service.get_phasing_view(scenario_id, fy_start_month=1)
    buckets = phasing.get("monthly", [])
    
    headers = ["Month", "Adds", "Removals", "Net HC", "Cost Delta", "Cum. Cost Delta"]
    rows_svg = []
    y_cursor = 100
    
    header_svg = f"""
      <rect x="50" y="{y_cursor}" width="1119" height="40" fill="#01244a" rx="4" ry="4"/>
      <text x="150" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Month</text>
      <text x="330" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Adds</text>
      <text x="510" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Removals</text>
      <text x="690" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Net HC</text>
      <text x="870" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Cost Delta</text>
      <text x="1050" y="{y_cursor + 20}" font-size="11" font-weight="700" fill="#ffffff" dominant-baseline="central" text-anchor="middle">Cum. Cost Delta</text>
    """
    rows_svg.append(header_svg)
    y_cursor += 40
    
    for i, b in enumerate(buckets):
        month = b["label"]
        adds = f"{b['adds_count']}"
        rem = f"{b['removes_count']}"
        net_hc = f"{b['net_hc_delta']:+d}"
        c_delta = f"${b['cost_delta']:+,.0f}"
        cum_cost = f"${b['cumulative_cost_delta']:+,.0f}"
        
        bg_color = "#ffffff" if i % 2 == 0 else "#f8f9fa"
        
        row_content = f"""
          <rect x="50" y="{y_cursor}" width="1119" height="35" fill="{bg_color}" stroke="#dce4ee" stroke-width="0.5"/>
          <text x="150" y="{y_cursor + 17.5}" font-size="10" fill="#4f6077" dominant-baseline="central" text-anchor="middle">{escape(month)}</text>
          <text x="330" y="{y_cursor + 17.5}" font-size="10" fill="#4f6077" dominant-baseline="central" text-anchor="middle">{adds}</text>
          <text x="510" y="{y_cursor + 17.5}" font-size="10" fill="#4f6077" dominant-baseline="central" text-anchor="middle">{rem}</text>
          <text x="690" y="{y_cursor + 17.5}" font-size="10" fill="#01244a" font-weight="600" dominant-baseline="central" text-anchor="middle">{net_hc}</text>
          <text x="870" y="{y_cursor + 17.5}" font-size="10" fill="{ '#27ae60' if b['cost_delta'] < 0 else '#c0392b' if b['cost_delta'] > 0 else '#4f6077' }" font-weight="600" dominant-baseline="central" text-anchor="middle">{c_delta}</text>
          <text x="1050" y="{y_cursor + 17.5}" font-size="10" fill="{ '#27ae60' if b['cumulative_cost_delta'] < 0 else '#c0392b' if b['cumulative_cost_delta'] > 0 else '#4f6077' }" font-weight="600" dominant-baseline="central" text-anchor="middle">{cum_cost}</text>
        """
        rows_svg.append(row_content)
        y_cursor += 35
        
    svg = f"""<svg xmlns="http://www.w3.org/2000/svg" width="1219" height="685" viewBox="0 0 1219 685" font-family="Inter, Arial, sans-serif">
      <rect x="0" y="0" width="1219" height="685" fill="#f4f6f9"/>
      <rect x="0" y="0" width="1219" height="60" fill="#01244a"/>
      <text x="30" y="35" font-size="20" font-weight="700" fill="#ffffff" dominant-baseline="central">Implementation Phasing Timeline</text>
      {"".join(rows_svg)}
    </svg>"""
    return svg


@router.get("/db/scenarios/{scenario_id}/export/pdf")
def db_export_scenario_pdf(
    scenario_id: int,
    project_id: int,
    detail: str = Query("summary", regex="^(overview|summary|full)$"),
    _user: dict = Depends(require_project_access()),
):
    scenario, dataset = _require_scenario_in_project(scenario_id, project_id)
    records = db_service.get_scenario_records(scenario_id)
    summary_data = db_service.get_scenario_summary(scenario_id)

    emp_col = dataset["emp_col"]
    mgr_col = dataset["mgr_col"]
    svg_kwargs = dict(
        emp_col=emp_col, mgr_col=mgr_col,
        job_title_col=dataset.get("job_title_col"),
        fte_col=dataset.get("fte_col"),
        flc_col=dataset.get("flc_col"),
        country_col=dataset.get("country_col"),
    )

    svg_pages: List[str] = []

    # 1. Executive Summary Table Page
    svg_pages.append(_render_summary_table_svg(summary_data, dataset, scenario))

    # 2. Scenario Comparison Page
    svg_pages.append(_render_comparison_table_svg(dataset["id"], scenario["id"], scenario["name"]))

    # 3. Phasing Page
    svg_pages.append(_render_phasing_table_svg(scenario["id"], scenario["name"]))

    # 4. Overview L1-L2 Page
    overview_svg = render_scenario_svg(
        records, **svg_kwargs,
        title=f"OrgSight 2.0  —  {scenario['name']}",
        subtitle=f"{dataset['name']}  ·  Overview (L1–L2)",
        max_depth=2,
    )
    svg_pages.append(overview_svg)

    if detail == "full":
        tree = get_tree_structure(records, emp_col, mgr_col)
        for root_id in tree["roots"]:
            for kid_id in tree["children"].get(root_id, []):
                rec = tree["by_id"].get(kid_id, {})
                jtc = dataset.get("job_title_col")
                kid_title = str(rec.get(jtc) or rec.get("Job Title") or kid_id) if jtc else str(rec.get("Job Title") or kid_id)
                kid_hc = tree["headcount"].get(kid_id, 0)
                if kid_hc < 1:
                    continue
                subtree_svg = render_scenario_svg(
                    records, **svg_kwargs,
                    title=f"{kid_title} — Team Structure",
                    subtitle=f"{kid_hc} headcount",
                    root_id=kid_id, max_depth=3,
                )
                svg_pages.append(subtree_svg)

    pdf_parts = []
    for svg_str in svg_pages:
        pdf_parts.append(_svg_to_pdf_bytes(svg_str))

    if len(pdf_parts) == 1:
        pdf_bytes = pdf_parts[0]
    else:
        from pypdf import PdfWriter
        merger = PdfWriter()
        for part in pdf_parts:
            merger.append(BytesIO(part))
        merged = BytesIO()
        merger.write(merged)
        merger.close()
        pdf_bytes = merged.getvalue()

    safe_name = "".join(c for c in scenario["name"] if c.isalnum() or c in "-_") or "scenario"
    return StreamingResponse(
        BytesIO(pdf_bytes), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="orgsight_{safe_name}.pdf"'},
    )


@router.get("/db/scenarios/{scenario_id}/export/records")
def db_export_scenario_records(
    scenario_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    scenario, _ = _require_scenario_in_project(scenario_id, project_id)
    records = db_service.get_scenario_records(scenario_id)
    if not records:
        df = pd.DataFrame()
    else:
        cleaned = [{k: v for k, v in r.items() if not k.startswith("__")} for r in records]
        df = pd.DataFrame(cleaned)

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Org Records")
    buf.seek(0)

    safe_name = "".join(c for c in scenario["name"] if c.isalnum() or c in "-_") or "scenario"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="orgsight_records_{safe_name}.xlsx"'},
    )


# ===========================================================================
# Activity Analysis Endpoints
# ===========================================================================

class ActivityConfigBody(BaseModel):
    name: str
    role_grouping_col: str
    dataset_id: int


class ActivitiesBody(BaseModel):
    activities: List[Dict[str, Any]]


class AllocationsBody(BaseModel):
    allocations: List[Dict[str, Any]]


class LeversBody(BaseModel):
    levers: List[Dict[str, Any]]


def _require_activity_config_in_project(config_id: int, project_id: int):
    from services import db_service
    cfg = db_service.get_activity_config(config_id)
    if not cfg:
        raise HTTPException(status_code=404, detail="Activity config not found")
    ds = db_service.get_dataset(cfg["dataset_id"])
    if not ds or ds.get("project_id") != project_id:
        raise HTTPException(status_code=403, detail="Activity config not in this project")
    return cfg


@router.post("/activity/configs")
def create_activity_config(
    body: ActivityConfigBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service
    _require_dataset_in_project(body.dataset_id, project_id)
    cfg = db_service.create_activity_config(body.dataset_id, body.name, body.role_grouping_col)
    return {"config": cfg}


@router.get("/activity/configs")
def list_activity_configs(
    dataset_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service
    _require_dataset_in_project(dataset_id, project_id)
    configs = db_service.get_activity_configs(dataset_id)
    return {"configs": configs}


@router.get("/activity/configs/{config_id}")
def get_activity_config(
    config_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    cfg = _require_activity_config_in_project(config_id, project_id)
    return {"config": cfg}


@router.get("/activity/configs/{config_id}/roles")
def get_activity_roles(
    config_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service, activity_service
    cfg = _require_activity_config_in_project(config_id, project_id)
    roles = activity_service.get_role_values(cfg["dataset_id"], cfg["role_grouping_col"])
    return {"roles": roles}


@router.post("/activity/configs/{config_id}/activities")
def upsert_activities(
    config_id: int,
    body: ActivitiesBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service
    _require_activity_config_in_project(config_id, project_id)
    activities = db_service.upsert_activities(config_id, body.activities)
    return {"activities": activities}


@router.post("/activity/configs/{config_id}/activities/upload")
async def upload_activities(
    config_id: int,
    project_id: int,
    file: UploadFile = File(...),
    _user: dict = Depends(require_project_access()),
):
    from services import db_service
    _require_activity_config_in_project(config_id, project_id)
    content = await file.read()
    df = pd.read_excel(BytesIO(content))
    cols = {str(c).strip().lower(): c for c in df.columns}
    name_col = next((cols[k] for k in ("activity", "activity name", "name") if k in cols), None)
    process_col = next((cols[k] for k in ("process", "process name", "category") if k in cols), None)
    desc_col = next((cols[k] for k in ("description", "desc") if k in cols), None)
    if not name_col:
        raise HTTPException(status_code=400, detail="Upload needs an 'Activity' or 'Name' column")
    activities_list = []
    for idx, row in df.iterrows():
        name = str(row[name_col]).strip()
        if not name or name.lower() in ("nan", "activity"):
            continue
        activities_list.append({
            "name": name,
            "process_name": str(row[process_col]).strip() if process_col else "",
            "description": str(row[desc_col]).strip() if desc_col else "",
        })
    activities = db_service.upsert_activities(config_id, activities_list)
    return {"activities": activities}


@router.get("/activity/configs/{config_id}/allocations")
def get_allocations(
    config_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service, activity_service
    _require_activity_config_in_project(config_id, project_id)
    matrix = db_service.get_role_allocation_matrix(config_id)
    completeness = activity_service.get_allocation_completeness(config_id)
    return {"matrix": matrix, "completeness": completeness}


@router.post("/activity/configs/{config_id}/allocations")
def upsert_allocations(
    config_id: int,
    body: AllocationsBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service
    _require_activity_config_in_project(config_id, project_id)
    db_service.upsert_role_allocations(config_id, body.allocations)
    matrix = db_service.get_role_allocation_matrix(config_id)
    return {"matrix": matrix}


@router.post("/activity/configs/{config_id}/allocations/upload")
async def upload_allocations(
    config_id: int,
    project_id: int,
    file: UploadFile = File(...),
    _user: dict = Depends(require_project_access()),
):
    from services import db_service, activity_service
    cfg = _require_activity_config_in_project(config_id, project_id)
    content = await file.read()
    df = pd.read_excel(BytesIO(content))
    activities = cfg.get("activities") or []
    if not activities:
        raise HTTPException(status_code=400, detail="Define activities before uploading allocations")
    allocs = activity_service.parse_allocation_upload(df, activities)
    if not allocs:
        raise HTTPException(status_code=400, detail="No valid allocations found in upload")
    db_service.upsert_role_allocations(config_id, allocs)
    matrix = db_service.get_role_allocation_matrix(config_id)
    return {"matrix": matrix, "count": len(allocs)}


@router.get("/activity/configs/{config_id}/levers")
def get_levers(
    config_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service
    _require_activity_config_in_project(config_id, project_id)
    levers = db_service.get_activity_levers(config_id)
    return {"levers": levers}


@router.post("/activity/configs/{config_id}/levers")
def upsert_levers(
    config_id: int,
    body: LeversBody,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service
    _require_activity_config_in_project(config_id, project_id)
    db_service.upsert_activity_levers(config_id, body.levers)
    return {"levers": db_service.get_activity_levers(config_id)}


@router.delete("/activity/configs/{config_id}/levers/{lever_id}")
def delete_lever(
    config_id: int,
    lever_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import db_service
    _require_activity_config_in_project(config_id, project_id)
    db_service.delete_activity_lever(lever_id)
    return {"status": "deleted"}


@router.post("/activity/configs/{config_id}/levers/upload")
async def upload_levers(
    config_id: int,
    project_id: int,
    file: UploadFile = File(...),
    _user: dict = Depends(require_project_access()),
):
    from services import db_service, activity_service
    cfg = _require_activity_config_in_project(config_id, project_id)
    content = await file.read()
    df = pd.read_excel(BytesIO(content))
    activities = cfg.get("activities") or []
    if not activities:
        raise HTTPException(status_code=400, detail="Define activities before uploading levers")
    try:
        levers = activity_service.parse_lever_upload(df, activities)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    if not levers:
        raise HTTPException(status_code=400, detail="No valid levers found in upload")
    db_service.upsert_activity_levers(config_id, levers)
    return {"levers": db_service.get_activity_levers(config_id), "count": len(levers)}


@router.post("/activity/configs/{config_id}/compute")
def compute_activity_impact(
    config_id: int,
    project_id: int,
    function_col: Optional[str] = Query(None),
    _user: dict = Depends(require_project_access()),
):
    from services import activity_service
    cfg = _require_activity_config_in_project(config_id, project_id)
    try:
        impact = activity_service.compute_impact(
            config_id, cfg["dataset_id"],
            function_col=function_col or None,
        )
        return {"impact": impact}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/activity/configs/{config_id}/impact/export")
def export_activity_impact(
    config_id: int,
    project_id: int,
    _user: dict = Depends(require_project_access()),
):
    from services import activity_service
    cfg = _require_activity_config_in_project(config_id, project_id)
    try:
        impact = activity_service.compute_impact(config_id, cfg["dataset_id"])
        excel_bytes = activity_service.impact_to_excel(impact, cfg["name"])
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    safe_name = "".join(c for c in cfg["name"] if c.isalnum() or c in "-_ ").strip().replace(" ", "_") or "activity"
    return StreamingResponse(
        BytesIO(excel_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="orgsight_activity_{safe_name}.xlsx"'},
    )
