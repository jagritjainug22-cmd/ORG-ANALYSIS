"""
Benchmark report export to a formatted Excel workbook.

Mirrors the on-screen report: a cover with the headline numbers, the narrative
sections, the full variance table with conditional colouring, the ranked
opportunities, and a methodology sheet that carries the provenance and coverage
caveats so an exported file cannot be circulated without them.
"""

from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services import benchmark_registry as reg

_NAVY = "01244A"
_BLUE = "0A3F86"
_LIGHT = "E8EEF5"
_GOOD = "E7F7EF"
_BAD = "FDECEE"
_NEUTRAL = "F5F7FA"
_AMBER_BG = "FFF7E6"

_TITLE = Font(bold=True, size=18, color=_NAVY)
_H2 = Font(bold=True, size=13, color=_NAVY)
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
_HEADER_FILL = PatternFill("solid", fgColor=_BLUE)
_THIN = Side(style="thin", color="D5DEE9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_VERDICT_FILL = {
    "favourable": PatternFill("solid", fgColor=_GOOD),
    "unfavourable": PatternFill("solid", fgColor=_BAD),
    "in_line": PatternFill("solid", fgColor=_NEUTRAL),
}


def build_workbook(report: Dict[str, Any], title: Optional[str] = None) -> bytes:
    wb = Workbook()
    _cover(wb.active, report, title)
    _narrative(wb.create_sheet("Analysis"), report)
    _variance(wb.create_sheet("Variance Detail"), report)
    _opportunities(wb.create_sheet("Opportunities"), report)
    _methodology(wb.create_sheet("Methodology"), report)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _header_row(ws, row: int, headers: List[str], widths: List[int]) -> None:
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _money(value: Optional[float], currency: str) -> str:
    if value in (None, 0):
        return "-"
    magnitude = abs(value)
    if magnitude >= 1_000_000:
        return f"{currency} {value / 1_000_000:,.2f}m"
    if magnitude >= 1_000:
        return f"{currency} {value / 1_000:,.0f}k"
    return f"{currency} {value:,.0f}"


def _cover(ws, report: Dict[str, Any], title: Optional[str]) -> None:
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 46

    summary = report.get("summary") or {}
    pack = report.get("pack") or {}
    coverage = report.get("coverage") or {}
    currency = summary.get("currency") or "USD"

    ws["A1"] = title or "OrgSight Benchmark Report"
    ws["A1"].font = _TITLE
    ws["A2"] = f"Generated {datetime.utcnow().strftime('%d %B %Y')}"
    ws["A2"].font = Font(size=10, color="64748B")

    rows = [
        ("Benchmark pack", pack.get("name"), pack.get("industry")),
        ("Effective year", pack.get("effective_year"), pack.get("region")),
        ("", "", ""),
        ("Headcount", summary.get("headcount"), ""),
        ("Total FTE", summary.get("total_fte"), ""),
        ("Total workforce cost", _money(summary.get("total_cost"), currency), ""),
        ("", "", ""),
        ("Health score", f"{summary.get('health_score')}/100", summary.get("health_band")),
        ("Comparisons made", summary.get("metrics_compared"), ""),
        ("Unfavourable / in line / favourable",
         f"{summary.get('unfavourable_count')} / {summary.get('in_line_count')} / "
         f"{summary.get('favourable_count')}", ""),
        ("Benchmark coverage", f"{coverage.get('coverage_pct')}% of FTE", ""),
    ]

    row = 4
    for label, value, extra in rows:
        if not label:
            row += 1
            continue
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color=_NAVY)
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=_LIGHT)
        ws.cell(row=row, column=2, value=value)
        ws.cell(row=row, column=3, value=extra)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Quantified opportunity").font = _H2
    row += 1
    realization = (f"{int((summary.get('realization_low') or 0.6) * 100)}-"
                   f"{int((summary.get('realization_high') or 0.7) * 100)}% realisation")
    _header_row(ws, row, ["Lens", "Opportunities", "Excess FTE", f"Realisable ({realization})"],
                [38, 30, 30, 46])
    ws.freeze_panes = None
    row += 1

    for key, label in (("functional", "Functional right-sizing"),
                       ("structural", "Structural / management")):
        bucket = summary.get(key) or {}
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=bucket.get("opportunity_count"))
        ws.cell(row=row, column=3, value=round(bucket.get("fte_gap") or 0, 1))
        ws.cell(row=row, column=4,
                value=f"{_money(bucket.get('savings_low'), currency)} - "
                      f"{_money(bucket.get('savings_high'), currency)}")
        for col in range(1, 5):
            ws.cell(row=row, column=col).border = _BORDER
        row += 1

    row += 1
    warning = ws.cell(row=row, column=1, value=summary.get("overlap_note"))
    warning.fill = PatternFill("solid", fgColor=_AMBER_BG)
    warning.font = Font(bold=True, color="8A5A00")
    warning.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=4)


_MD_STRIP = re.compile(r"(\*\*|__|`)")


def _narrative(ws, report: Dict[str, Any]) -> None:
    ws.column_dimensions["A"].width = 120
    row = 1
    for section in report.get("sections") or []:
        ws.cell(row=row, column=1, value=section.get("title")).font = _H2
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=_LIGHT)
        row += 2
        for line in (section.get("markdown") or "").split("\n"):
            text = _MD_STRIP.sub("", line).rstrip()
            if not text:
                row += 1
                continue
            cell = ws.cell(row=row, column=1, value=text.lstrip("#").strip()
                           if text.lstrip().startswith("#") else text)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if text.lstrip().startswith("#"):
                cell.font = Font(bold=True, size=11, color=_BLUE)
            row += 1
        row += 2


def _variance(ws, report: Dict[str, Any]) -> None:
    currency = (report.get("summary") or {}).get("currency") or "USD"
    headers = ["Scope", "Function", "Sub-function", "Metric", "Client", "P25",
               "Median", "P75", "Variance %", "Position", "Verdict",
               "FTE gap", "Cost gap"]
    _header_row(ws, 1, headers, [13, 24, 34, 32, 14, 12, 13, 12, 13, 14, 15, 12, 16])

    rows = sorted(
        [r for r in report.get("variance") or [] if r.get("client_value") is not None],
        key=lambda r: (r["scope"], -(abs(r.get("delta_pct") or 0))),
    )
    for idx, r in enumerate(rows, start=2):
        key = r["metric_key"]
        values = [
            r["scope"], r.get("function"), r.get("subfunction"), r.get("metric_label"),
            reg.format_value(key, r.get("client_value"), currency),
            reg.format_value(key, r.get("p25"), currency),
            reg.format_value(key, r.get("median"), currency),
            reg.format_value(key, r.get("p75"), currency),
            None if r.get("delta_pct") is None else round(r["delta_pct"], 1),
            (r.get("position") or "").replace("_", " "),
            (r.get("verdict") or "").replace("_", " "),
            None if r.get("fte_gap") is None else round(r["fte_gap"], 1),
            _money(r.get("cost_gap"), currency) if r.get("cost_gap") else "-",
        ]
        fill = _VERDICT_FILL.get(r.get("verdict"))
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=idx, column=col, value=value)
            cell.border = _BORDER
            if fill:
                cell.fill = fill
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"


def _opportunities(ws, report: Dict[str, Any]) -> None:
    currency = (report.get("summary") or {}).get("currency") or "USD"
    headers = ["Rank", "Lens", "Area", "Primary metric", "Client", "Median",
               "Variance %", "Excess FTE", "Gross cost gap", "Realisable low",
               "Realisable high", "Confidence"]
    _header_row(ws, 1, headers, [8, 15, 30, 32, 14, 14, 12, 12, 18, 18, 18, 13])

    for idx, o in enumerate(report.get("opportunities") or [], start=2):
        key = o.get("primary_metric")
        values = [
            o.get("rank"), o.get("bucket"), o.get("label"), o.get("primary_metric_label"),
            reg.format_value(key, o.get("client_value"), currency),
            reg.format_value(key, o.get("median"), currency),
            None if o.get("delta_pct") is None else round(o["delta_pct"], 1),
            None if o.get("fte_gap") is None else round(o["fte_gap"], 1),
            _money(o.get("cost_gap"), currency),
            _money(o.get("savings_low"), currency),
            _money(o.get("savings_high"), currency),
            o.get("confidence"),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=idx, column=col, value=value)
            cell.border = _BORDER
            if col == 2:
                cell.fill = PatternFill(
                    "solid", fgColor=_LIGHT if o.get("bucket") == "functional" else _AMBER_BG)


def _methodology(ws, report: Dict[str, Any]) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100

    coverage = report.get("coverage") or {}
    pack = report.get("pack") or {}
    columns = report.get("columns") or {}

    ws["A1"] = "Methodology, provenance and limitations"
    ws["A1"].font = _TITLE

    row = 3
    entries: List[tuple] = [
        ("Benchmark pack", f"{pack.get('name')} — {pack.get('industry')}, {pack.get('region')}, "
                           f"{pack.get('effective_year')}"),
        ("Pack notes", pack.get("notes") or "-"),
        ("Provenance", "Built-in packs shipped with OrgSight are directional and must be validated "
                       "before client use. Custom packs carry the source recorded at upload."
         if pack.get("is_builtin") else "Custom pack uploaded by the engagement team."),
        ("Coverage", f"{coverage.get('coverage_pct')}% of FTE maps to a benchmarked function "
                     f"({coverage.get('mapped_fte')} of {coverage.get('total_fte')} FTE)."),
        ("Unmapped functions", ", ".join(coverage.get("unmapped_functions") or []) or "None"),
        ("Benchmark functions with no client match",
         ", ".join(coverage.get("benchmark_functions_absent_from_client") or []) or "None"),
        ("Function column", columns.get("func_col") or "not mapped"),
        ("Sub-function column", columns.get("subfunc_col") or "not mapped"),
        ("FTE column", columns.get("fte_col") or "not mapped — headcount used as a proxy"),
        ("Cost column", columns.get("flc_col") or "not mapped — cost metrics unavailable"),
        ("Savings basis", (report.get("summary") or {}).get("overlap_note")),
    ]

    for label, value in entries:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color=_NAVY)
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=_LIGHT)
        cell = ws.cell(row=row, column=2, value=value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Warnings").font = _H2
    row += 1
    warnings = list(coverage.get("warnings") or [])
    if (report.get("currency") or {}).get("message"):
        warnings.append(report["currency"]["message"])
    for s in report.get("suppressed_metrics") or []:
        if s.get("reason"):
            warnings.append(f"{s['label']}: {s['reason']}")
    for warning in warnings or ["None."]:
        cell = ws.cell(row=row, column=2, value=warning)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = PatternFill("solid", fgColor=_AMBER_BG)
        row += 1
