"""
Benchmark pack Excel template — generation and parsing.

The template is deliberately wide and human-friendly (one row per function,
one column per metric) because that is how consultants actually hold benchmark
data. The parser normalises it into the long metric rows the engine consumes,
so the storage format stays flexible while the authoring format stays simple.
"""

from __future__ import annotations

import json
import logging
import os
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from services import benchmark_registry as reg

log = logging.getLogger(__name__)

SHEET_README = "README"
SHEET_PACK_INFO = "Pack Info"
SHEET_FUNCTIONS = "Functional Benchmarks"
SHEET_SUBFUNCTIONS = "Subfunction Benchmarks"
SHEET_STRUCTURAL = "Structural Benchmarks"
SHEET_REFERENCE = "Reference Lists"

_NAVY = "01244A"
_BLUE = "0A3F86"
_LIGHT = "E8EEF5"
_AMBER = "FFF7E6"

_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill("solid", fgColor=_BLUE)
_TITLE_FONT = Font(bold=True, size=14, color=_NAVY)
_SUB_FILL = PatternFill("solid", fgColor=_LIGHT)
_NOTE_FILL = PatternFill("solid", fgColor=_AMBER)
_THIN = Side(style="thin", color="C9D6E4")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _taxonomy() -> Dict[str, List[str]]:
    path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "master_taxonomy.json")
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        log.warning("Could not load master taxonomy for the benchmark template: %s", e)
        return {}


# ---------------------------------------------------------------------------
# Generation
# ---------------------------------------------------------------------------

def build_template(prefill_pack: Optional[Dict[str, Any]] = None,
                   prefill_metrics: Optional[List[Dict[str, Any]]] = None) -> bytes:
    """Build the upload template, optionally pre-filled with an existing pack."""
    taxonomy = _taxonomy()
    wb = Workbook()

    _build_readme(wb.active)
    _build_pack_info(wb.create_sheet(SHEET_PACK_INFO), prefill_pack)
    _build_functional(wb.create_sheet(SHEET_FUNCTIONS), taxonomy, prefill_metrics)
    _build_subfunctional(wb.create_sheet(SHEET_SUBFUNCTIONS), taxonomy, prefill_metrics)
    _build_structural(wb.create_sheet(SHEET_STRUCTURAL), prefill_metrics)
    _build_reference(wb.create_sheet(SHEET_REFERENCE), taxonomy)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _style_header(ws, row: int, count: int) -> None:
    for col in range(1, count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[row].height = 34
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _build_readme(ws) -> None:
    ws.title = SHEET_README
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 70

    ws["A1"] = "OrgSight Benchmark Pack Template"
    ws["A1"].font = _TITLE_FONT

    instructions = [
        "",
        "How to use this template:",
        "1. Fill in 'Pack Info' — the pack name and currency are required.",
        "2. Fill in 'Functional Benchmarks' — one row per function. Leave any metric blank if you do not have it.",
        "3. Optionally fill 'Subfunction Benchmarks' and 'Structural Benchmarks'.",
        "4. Every metric column accepts a median value. The matching '(P25)' and '(P75)' columns are optional but",
        "   strongly recommended: without a range, OrgSight falls back to a +/-10% band around the median.",
        "5. Use the function names on the 'Reference Lists' sheet wherever possible. They match the A&M master",
        "   taxonomy that OrgSight rationalises client data onto, so the comparison joins cleanly.",
        "6. Upload the completed file in the Benchmarking tab.",
        "",
        "Metric dictionary:",
    ]
    row = 2
    for line in instructions:
        ws.cell(row=row, column=1, value=line)
        if line.endswith(":"):
            ws.cell(row=row, column=1).font = Font(bold=True, color=_NAVY)
        row += 1

    headers = ["Metric", "Applies to", "Unit", "Direction", "Needs context", "What it measures"]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=row, column=idx, value=header)
    _style_header(ws, row, len(headers))
    ws.freeze_panes = None
    row += 1

    for meta in reg.metric_dictionary():
        ws.cell(row=row, column=1, value=meta["label"])
        ws.cell(row=row, column=2, value=", ".join(meta["scopes"]))
        ws.cell(row=row, column=3, value=meta["unit"])
        ws.cell(row=row, column=4, value=meta["direction"].replace("_", " "))
        ws.cell(row=row, column=5, value=", ".join(meta["requires_context"]) or "-")
        cell = ws.cell(row=row, column=6, value=meta["description"])
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1


def _build_pack_info(ws, prefill: Optional[Dict[str, Any]]) -> None:
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 60

    ws["A1"] = "Field"
    ws["B1"] = "Value"
    ws["C1"] = "Guidance"
    _style_header(ws, 1, 3)

    p = prefill or {}
    fields = [
        ("Pack name", p.get("name") or "", "Required. Shown in the pack picker, e.g. 'Client X Peer Set 2026'."),
        ("Industry", p.get("industry") or "", "e.g. Financial Services, Manufacturing & Industrials."),
        ("Region", p.get("region") or "Global", "e.g. Global, EMEA, North America."),
        ("Size band", p.get("size_band") or "All sizes", "e.g. Under 1,000 / 1,000-10,000 / 10,000+."),
        ("Currency", p.get("currency") or "USD", "ISO code for any absolute cost values in this pack."),
        ("Effective year", p.get("effective_year") or "", "Year the benchmark data represents."),
        ("Source", "", "Where the data came from. Recorded against every row for auditability."),
        ("Notes", p.get("notes") or "", "Caveats, peer set composition, exclusions."),
    ]
    for idx, (field, value, guidance) in enumerate(fields, start=2):
        ws.cell(row=idx, column=1, value=field).font = Font(bold=True)
        ws.cell(row=idx, column=1).fill = _SUB_FILL
        ws.cell(row=idx, column=2, value=value)
        ws.cell(row=idx, column=3, value=guidance).alignment = Alignment(wrap_text=True)
        for col in range(1, 4):
            ws.cell(row=idx, column=col).border = _BORDER


def _metric_columns(pairs: List[Tuple[str, str]]) -> List[Tuple[str, str, Optional[str]]]:
    """Expand each metric into median, P25 and P75 spreadsheet columns."""
    out: List[Tuple[str, str, Optional[str]]] = []
    for header, metric_key in pairs:
        out.append((header, metric_key, "median"))
        out.append((f"{header} (P25)", metric_key, "p25"))
        out.append((f"{header} (P75)", metric_key, "p75"))
    return out


def _index_prefill(metrics: Optional[List[Dict[str, Any]]]) -> Dict[Tuple[Any, ...], Dict[str, Any]]:
    index: Dict[Tuple[Any, ...], Dict[str, Any]] = {}
    for m in metrics or []:
        index[(m.get("scope"), m.get("function"), m.get("subfunction"), m.get("metric_key"))] = m
    return index


def _build_functional(ws, taxonomy: Dict[str, List[str]],
                      prefill: Optional[List[Dict[str, Any]]]) -> None:
    columns = _metric_columns(reg.FUNCTIONAL_TEMPLATE_COLUMNS)
    headers = ["Function"] + [c[0] for c in columns]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=header)
    _style_header(ws, 1, len(headers))
    ws.column_dimensions["A"].width = 28
    for idx in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 15

    index = _index_prefill(prefill)
    functions = list(taxonomy.keys()) or ["Finance", "HR", "IT"]
    for r, function in enumerate(functions, start=2):
        ws.cell(row=r, column=1, value=function).font = Font(bold=True)
        ws.cell(row=r, column=1).fill = _SUB_FILL
        for c, (_, metric_key, stat) in enumerate(columns, start=2):
            row_data = index.get(("function", function, None, metric_key))
            cell = ws.cell(row=r, column=c, value=(row_data or {}).get(stat))
            cell.border = _BORDER

    _attach_function_validation(ws, taxonomy, column=1, last_row=len(functions) + 40)


def _build_subfunctional(ws, taxonomy: Dict[str, List[str]],
                         prefill: Optional[List[Dict[str, Any]]]) -> None:
    columns = _metric_columns(reg.SUBFUNCTIONAL_TEMPLATE_COLUMNS)
    headers = ["Function", "Subfunction"] + [c[0] for c in columns]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=header)
    _style_header(ws, 1, len(headers))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 56
    for idx in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 17

    note = ws.cell(row=1, column=len(headers) + 2,
                   value="Shares are measured against the PARENT FUNCTION, not the whole organisation.")
    note.fill = _NOTE_FILL
    note.font = Font(bold=True, color="8A5A00")

    index = _index_prefill(prefill)
    r = 2
    for function, subfunctions in (taxonomy or {}).items():
        for subfunction in subfunctions:
            ws.cell(row=r, column=1, value=function)
            ws.cell(row=r, column=2, value=subfunction)
            for c, (_, metric_key, stat) in enumerate(columns, start=3):
                row_data = index.get(("subfunction", function, subfunction, metric_key))
                cell = ws.cell(row=r, column=c, value=(row_data or {}).get(stat))
                cell.border = _BORDER
            r += 1


def _build_structural(ws, prefill: Optional[List[Dict[str, Any]]]) -> None:
    headers = ["Metric", "Metric key", "Unit", "P25", "Median", "P75"]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=header)
    _style_header(ws, 1, len(headers))
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 12
    for col in ("D", "E", "F"):
        ws.column_dimensions[col].width = 13

    index = _index_prefill(prefill)
    for r, metric_key in enumerate(reg.STRUCTURAL_TEMPLATE_METRICS, start=2):
        meta = reg.get_metric(metric_key)
        if not meta:
            continue
        ws.cell(row=r, column=1, value=meta["label"]).font = Font(bold=True)
        key_cell = ws.cell(row=r, column=2, value=metric_key)
        key_cell.font = Font(color="7A8798", size=9)
        ws.cell(row=r, column=3, value=meta["unit"])
        row_data = index.get(("org", None, None, metric_key)) or {}
        ws.cell(row=r, column=4, value=row_data.get("p25"))
        ws.cell(row=r, column=5, value=row_data.get("median"))
        ws.cell(row=r, column=6, value=row_data.get("p75"))
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = _BORDER


def _build_reference(ws, taxonomy: Dict[str, List[str]]) -> None:
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 66
    ws["A1"] = "Function"
    ws["B1"] = "Subfunction"
    _style_header(ws, 1, 2)

    r = 2
    for function, subfunctions in (taxonomy or {}).items():
        for subfunction in subfunctions:
            ws.cell(row=r, column=1, value=function)
            ws.cell(row=r, column=2, value=subfunction)
            r += 1

    ws.cell(row=r + 1, column=1, value="These are the A&M master taxonomy values that OrgSight "
                                       "rationalises client function data onto. Using them makes the "
                                       "benchmark join exact.").fill = _NOTE_FILL


def _attach_function_validation(ws, taxonomy: Dict[str, List[str]], column: int, last_row: int) -> None:
    functions = list(taxonomy.keys())
    if not functions:
        return
    formula = '"' + ",".join(functions) + '"'
    if len(formula) > 255:  # Excel inline-list limit
        return
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    letter = get_column_letter(column)
    dv.add(f"{letter}2:{letter}{max(last_row, 60)}")


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------

class TemplateError(ValueError):
    """Raised when an uploaded workbook cannot be interpreted."""


def parse_template(content: bytes) -> Dict[str, Any]:
    """Parse an uploaded workbook into pack metadata plus long metric rows."""
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as e:
        raise TemplateError(f"Could not read the workbook: {e}") from e

    info = _parse_pack_info(wb)
    source = info.pop("_source", None) or "Custom upload"

    metrics: List[Dict[str, Any]] = []
    warnings: List[str] = []

    metrics += _parse_wide_sheet(
        wb, SHEET_FUNCTIONS, reg.FUNCTIONAL_TEMPLATE_COLUMNS,
        scope="function", key_columns=["Function"], source=source, warnings=warnings,
    )
    metrics += _parse_wide_sheet(
        wb, SHEET_SUBFUNCTIONS, reg.SUBFUNCTIONAL_TEMPLATE_COLUMNS,
        scope="subfunction", key_columns=["Function", "Subfunction"],
        source=source, warnings=warnings,
    )
    metrics += _parse_structural(wb, source, warnings)

    if not metrics:
        raise TemplateError(
            "No benchmark values were found. Fill in at least one metric on the "
            "'Functional Benchmarks' or 'Structural Benchmarks' sheet."
        )
    if not info.get("name"):
        raise TemplateError("'Pack name' is required on the 'Pack Info' sheet.")

    return {"pack": info, "metrics": metrics, "warnings": warnings}


def _parse_pack_info(wb) -> Dict[str, Any]:
    if SHEET_PACK_INFO not in wb.sheetnames:
        raise TemplateError(f"Missing the '{SHEET_PACK_INFO}' sheet. Download a fresh template.")
    ws = wb[SHEET_PACK_INFO]
    values: Dict[str, Any] = {}
    for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        label, value = (row + (None, None))[:2]
        if label is None:
            continue
        values[str(label).strip().lower()] = value

    year = values.get("effective year")
    try:
        year = int(year) if year not in (None, "") else None
    except (TypeError, ValueError):
        year = None

    return {
        "name": _clean(values.get("pack name")),
        "industry": _clean(values.get("industry")),
        "region": _clean(values.get("region")),
        "size_band": _clean(values.get("size band")),
        "currency": (_clean(values.get("currency")) or "USD").upper(),
        "effective_year": year,
        "notes": _clean(values.get("notes")),
        "_source": _clean(values.get("source")),
    }


def _clean(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _header_map(ws) -> Dict[str, int]:
    headers: Dict[str, int] = {}
    for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1):
        if cell.value is not None:
            headers[str(cell.value).strip().lower()] = idx
    return headers


def _parse_wide_sheet(
    wb, sheet_name: str, column_pairs: List[Tuple[str, str]], scope: str,
    key_columns: List[str], source: str, warnings: List[str],
) -> List[Dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    headers = _header_map(ws)

    missing_keys = [k for k in key_columns if k.lower() not in headers]
    if missing_keys:
        warnings.append(f"'{sheet_name}' is missing the {', '.join(missing_keys)} column and was skipped.")
        return []

    out: List[Dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2):
        cells = {idx + 1: c.value for idx, c in enumerate(row)}
        key_values = [_clean(cells.get(headers[k.lower()])) for k in key_columns]
        if not key_values[0]:
            continue
        function = key_values[0]
        subfunction = key_values[1] if len(key_values) > 1 else None
        if scope == "subfunction" and not subfunction:
            continue

        for header, metric_key in column_pairs:
            meta = reg.get_metric(metric_key)
            if not meta:
                continue
            stats = {
                "median": _number(cells.get(headers.get(header.lower()))),
                "p25": _number(cells.get(headers.get(f"{header} (p25)".lower()))),
                "p75": _number(cells.get(headers.get(f"{header} (p75)".lower()))),
            }
            if all(v is None for v in stats.values()):
                continue
            if stats["median"] is None:
                bounds = [v for v in (stats["p25"], stats["p75"]) if v is not None]
                stats["median"] = sum(bounds) / len(bounds)
            if stats["p25"] is not None and stats["p75"] is not None and stats["p25"] > stats["p75"]:
                stats["p25"], stats["p75"] = stats["p75"], stats["p25"]
                warnings.append(
                    f"{function}{' / ' + subfunction if subfunction else ''} — "
                    f"'{header}' had P25 above P75; the two were swapped."
                )

            out.append({
                "scope": scope,
                "function": function,
                "subfunction": subfunction,
                "metric_key": metric_key,
                "unit": meta["unit"],
                "p25": stats["p25"],
                "median": stats["median"],
                "p75": stats["p75"],
                "direction": meta["direction"],
                "source": source,
                "notes": None,
            })
    return out


def _parse_structural(wb, source: str, warnings: List[str]) -> List[Dict[str, Any]]:
    if SHEET_STRUCTURAL not in wb.sheetnames:
        return []
    ws = wb[SHEET_STRUCTURAL]
    headers = _header_map(ws)
    if "metric key" not in headers:
        warnings.append(f"'{SHEET_STRUCTURAL}' is missing the 'Metric key' column and was skipped.")
        return []

    label_to_key = {m["label"].lower(): m["metric_key"] for m in reg.metric_dictionary()}
    out: List[Dict[str, Any]] = []

    for row in ws.iter_rows(min_row=2):
        cells = {idx + 1: c.value for idx, c in enumerate(row)}
        metric_key = _clean(cells.get(headers["metric key"]))
        if not metric_key:
            label = _clean(cells.get(headers.get("metric", 0)))
            metric_key = label_to_key.get((label or "").lower())
        if not metric_key:
            continue
        meta = reg.get_metric(metric_key)
        if not meta:
            warnings.append(f"Unknown metric key '{metric_key}' on '{SHEET_STRUCTURAL}' was ignored.")
            continue

        p25 = _number(cells.get(headers.get("p25")))
        median = _number(cells.get(headers.get("median")))
        p75 = _number(cells.get(headers.get("p75")))
        if median is None and p25 is None and p75 is None:
            continue
        if median is None:
            bounds = [v for v in (p25, p75) if v is not None]
            median = sum(bounds) / len(bounds)

        out.append({
            "scope": "org",
            "function": None,
            "subfunction": None,
            "metric_key": metric_key,
            "unit": meta["unit"],
            "p25": p25,
            "median": median,
            "p75": p75,
            "direction": meta["direction"],
            "source": source,
            "notes": None,
        })
    return out


def _number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None
